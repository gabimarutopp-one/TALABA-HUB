import asyncio
import io
import json
import logging
import os
import re
import sqlite3
import tempfile
import threading
from datetime import datetime

import aiohttp
import cv2
import numpy as np
from cryptography.fernet import Fernet, InvalidToken
from aiogram import Bot, Dispatcher, F, BaseMiddleware
from aiogram.client.session.aiohttp import AiohttpSession
from aiogram.filters import CommandStart, Command, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (
    Message,
    BufferedInputFile,
    InlineKeyboardMarkup,
    InlineKeyboardButton,
    ReplyKeyboardMarkup,
    KeyboardButton,
    CallbackQuery,
    BotCommand,
    BotCommandScopeChat,
    BotCommandScopeDefault,
    TelegramObject,
    InputMediaDocument,
    WebAppInfo,
)
from PIL import Image, ImageOps, ImageFilter, ImageDraw
from rembg import remove, new_session
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

logging.basicConfig(level=logging.INFO)

# ====== SOZLAMALAR ======
BOT_TOKEN = os.getenv("BOT_TOKEN", "")  # @BotFather dan olinadi, muhit o'zgaruvchisidan o'qiladi
if not BOT_TOKEN:
    raise RuntimeError(
        "BOT_TOKEN muhit o'zgaruvchisi topilmadi. "
        "PythonAnywhere'da: Web sahifasida 'Environment variables' bo'limiga qo'shing, "
        "yoki lokal ishga tushirishda: export BOT_TOKEN=... (Linux/Mac) yoki "
        "$env:BOT_TOKEN=\"...\" (Windows PowerShell)."
    )

# Admin(lar)ning Telegram user_id raqamlari.
# Muhit o'zgaruvchisi orqali berish mumkin: ADMIN_IDS="123456789,987654321"
# Yoki shu yerga to'g'ridan-to'g'ri yozib qo'yishingiz mumkin: {123456789}
ADMIN_IDS: set[int] = {6137921070} | {
    int(x) for x in os.getenv("ADMIN_IDS", "").split(",") if x.strip().isdigit()
}

DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "bot_stats.db")

# ====== TO'LOV SOZLAMALARI ======
# Xizmatlar narxi (so'mda). Adminlar bu narxlardan bepul foydalanadi.
PRICE_3X4 = 10_000          # 3x4 rasm narxi
PRICE_OBYEKTIVKA = 15_000   # Ma'lumotnoma (obyektivka) narxi

# TO'LOV UCHUN KARTA MA'LUMOTLARI
PAYMENT_CARD_NUMBER = "6262 5700 4746 8014"
PAYMENT_CARD_OWNER = "A.Ja'far"

# 3x4 sm, 300 DPI da piksel o'lchami (chop etish uchun sifatli)
TARGET_WIDTH = 354   # 3 sm
TARGET_HEIGHT = 472  # 4 sm

# Yakuniy rasmda "bosh" (peshonadan-iyakkacha) balandligi qancha ulush
# egallashi kerakligi.
HEAD_HEIGHT_RATIO = 0.62

# Rasm sifatini oldindan tekshirish uchun chegaralar
MIN_IMAGE_DIMENSION = 200      # piksel (kenglik yoki balandlik shundan kichik bo'lmasin)
BLUR_VARIANCE_THRESHOLD = 60.0  # Laplasian dispersiyasi shundan past bo'lsa - loyqa


# ============================================================
# SQLITE: FOYDALANUVCHILAR / SO'ROVLAR / XATOLIKLAR STATISTIKASI
# ============================================================

_db_lock = threading.Lock()


# ====== HEMIS SOZLAMALARI ======
# HEMIS har bir universitet uchun alohida domenda ishlaydi
# (masalan: student.tuit.uz, student.nuu.uz va h.k.).
# Foydalanuvchi o'z domenini /hemis buyrug'i orqali kiritadi.
#
# DIQQAT: quyidagi endpoint yo'llari ko'pchilik HEMIS instansiyalarida
# ishlaydigan eng keng tarqalgan andoza asosida yozilgan. Ba'zi
# universitetlarda yo'l yoki javob strukturasi biroz farq qilishi
# mumkin - shuning uchun o'z hisobingiz bilan sinab ko'ring va
# kerak bo'lsa shu konstantalarni moslashtiring.
HEMIS_LOGIN_PATH = "/auth/login"
HEMIS_GPA_PATH = "/education/gpa-list"
HEMIS_SUBJECTS_PATH = "/education/subject-list"
HEMIS_REQUEST_TIMEOUT = 15

# Talaba sessiyasi tugab qolganda, foydalanuvchi qayta login/parol
# yozmasligi uchun parol shifrlangan holda saqlanadi va token muddati
# o'tganda shu parol bilan avtomatik qayta login qilinadi.
# Kalitni generatsiya qilish: python -c "from cryptography.fernet import Fernet; print(Fernet.generate_key().decode())"
HEMIS_ENC_KEY = os.getenv("HEMIS_ENC_KEY")
if HEMIS_ENC_KEY:
    _fernet = Fernet(HEMIS_ENC_KEY.encode())
else:
    _fernet = None
    logging.warning(
        "HEMIS_ENC_KEY topilmadi - talaba paroli saqlanmaydi va token muddati "
        "tugaganda foydalanuvchi qayta ulanishi kerak bo'ladi."
    )


def _encrypt_text(text: str) -> str | None:
    if not _fernet or not text:
        return None
    return _fernet.encrypt(text.encode()).decode()


def _decrypt_text(token: str | None) -> str | None:
    if not _fernet or not token:
        return None
    try:
        return _fernet.decrypt(token.encode()).decode()
    except InvalidToken:
        return None

# HEMIS mini-app (Telegram WebApp) manzili. Bu domen PythonAnywhere (yoki
# boshqa https hostingga) joylashtirilgach to'ldiriladi.
HEMIS_WEBAPP_URL = os.getenv("HEMIS_WEBAPP_URL", "")

# HEMIS xizmatini vaqtincha o'chirib qo'yish uchun (masalan hosting geo-blok
# muammosi hal bo'lguncha). "true"/"false" - muhit o'zgaruvchisi orqali ham
# boshqarish mumkin: HEMIS_SERVICE_ENABLED=false
HEMIS_SERVICE_ENABLED = os.getenv("HEMIS_SERVICE_ENABLED", "false").strip().lower() == "true"
HEMIS_UNAVAILABLE_TEXT = "🛠 Bu xizmat vaqtincha ishlamayapti. Keyinroq qaytadan urinib ko'ring."


def db_init() -> None:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            """CREATE TABLE IF NOT EXISTS users (
                user_id INTEGER PRIMARY KEY,
                username TEXT,
                first_seen TEXT
            )"""
        )
        conn.execute(
            """CREATE TABLE IF NOT EXISTS requests (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                created_at TEXT
            )"""
        )
        conn.execute(
            """CREATE TABLE IF NOT EXISTS errors (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                error_text TEXT,
                created_at TEXT
            )"""
        )
        # Dinamik (kod ichida emas, komandalar orqali) qo'shiladigan adminlar
        conn.execute(
            """CREATE TABLE IF NOT EXISTS admins (
                user_id INTEGER PRIMARY KEY,
                added_by INTEGER,
                added_at TEXT
            )"""
        )
        # Foydalanuvchi yuborgan to'lov cheklari va ularning holati
        conn.execute(
            """CREATE TABLE IF NOT EXISTS pending_payments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                username TEXT,
                full_name TEXT,
                service TEXT,
                amount INTEGER,
                receipt_file_id TEXT,
                status TEXT NOT NULL DEFAULT 'pending',
                created_at TEXT,
                decided_at TEXT,
                decided_by INTEGER
            )"""
        )
        # Foydalanuvchining HEMIS hisobi: domen, login va joriy token.
        # Parol shifrlangan holda saqlanadi (HEMIS_ENC_KEY orqali) - shunda
        # token muddati tugaganda foydalanuvchi qayta yozmasdan avtomatik
        # qayta ulanadi.
        conn.execute(
            """CREATE TABLE IF NOT EXISTS hemis_accounts (
                user_id INTEGER PRIMARY KEY,
                domain TEXT,
                login TEXT,
                token TEXT,
                password_enc TEXT,
                updated_at TEXT
            )"""
        )
        # Eski bazalarda password_enc ustuni bo'lmasligi mumkin - qo'shamiz
        try:
            conn.execute("ALTER TABLE hemis_accounts ADD COLUMN password_enc TEXT")
        except sqlite3.OperationalError:
            pass  # ustun allaqachon mavjud
        conn.commit()


def db_add_user(user_id: int, username: str | None) -> None:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            "INSERT OR IGNORE INTO users (user_id, username, first_seen) VALUES (?, ?, ?)",
            (user_id, username or "", datetime.now().isoformat(timespec="seconds")),
        )
        conn.commit()


def db_log_request(user_id: int) -> None:
    """Har bir /malumotnoma (tayyor hujjat yuborilgan) so'rovini yozib boradi."""
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            "INSERT INTO requests (user_id, created_at) VALUES (?, ?)",
            (user_id, datetime.now().isoformat(timespec="seconds")),
        )
        conn.commit()


def db_log_error(user_id: int, error_text: str) -> None:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            "INSERT INTO errors (user_id, error_text, created_at) VALUES (?, ?, ?)",
            (user_id, str(error_text)[:500], datetime.now().isoformat(timespec="seconds")),
        )
        conn.commit()


def db_get_stats() -> dict:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        total_users = conn.execute("SELECT COUNT(*) FROM users").fetchone()[0]
        total_requests = conn.execute("SELECT COUNT(*) FROM requests").fetchone()[0]
        total_errors = conn.execute("SELECT COUNT(*) FROM errors").fetchone()[0]
        today = datetime.now().strftime("%Y-%m-%d")
        today_requests = conn.execute(
            "SELECT COUNT(*) FROM requests WHERE created_at LIKE ?", (f"{today}%",)
        ).fetchone()[0]
        return {
            "total_users": total_users,
            "total_requests": total_requests,
            "total_errors": total_errors,
            "today_requests": today_requests,
        }


def db_get_payment_stats() -> dict:
    """To'lovlar bo'yicha kengaytirilgan statistika: holat, xizmat va tushum."""
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        by_status = dict(
            conn.execute(
                "SELECT status, COUNT(*) FROM pending_payments GROUP BY status"
            ).fetchall()
        )
        by_service = conn.execute(
            "SELECT service, COUNT(*) FROM pending_payments "
            "WHERE status = 'confirmed' GROUP BY service"
        ).fetchall()
        revenue = conn.execute(
            "SELECT COALESCE(SUM(amount), 0) FROM pending_payments WHERE status = 'confirmed'"
        ).fetchone()[0]
        today = datetime.now().strftime("%Y-%m-%d")
        revenue_today = conn.execute(
            "SELECT COALESCE(SUM(amount), 0) FROM pending_payments "
            "WHERE status = 'confirmed' AND decided_at LIKE ?",
            (f"{today}%",),
        ).fetchone()[0]
        return {
            "pending": by_status.get("pending", 0),
            "confirmed": by_status.get("confirmed", 0),
            "rejected": by_status.get("rejected", 0),
            "by_service": dict(by_service),
            "revenue": revenue,
            "revenue_today": revenue_today,
        }


def db_get_all_user_ids() -> list[int]:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        rows = conn.execute("SELECT user_id FROM users").fetchall()
        return [row[0] for row in rows]


def db_set_hemis_account(
    user_id: int, domain: str, login: str, token: str, password: str | None = None
) -> None:
    password_enc = _encrypt_text(password) if password else None
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        if password_enc is not None:
            conn.execute(
                "INSERT INTO hemis_accounts (user_id, domain, login, token, password_enc, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?) "
                "ON CONFLICT(user_id) DO UPDATE SET "
                "domain=excluded.domain, login=excluded.login, token=excluded.token, "
                "password_enc=excluded.password_enc, updated_at=excluded.updated_at",
                (user_id, domain, login, token, password_enc, datetime.now().isoformat(timespec="seconds")),
            )
        else:
            # Parol berilmagan (masalan token yangilanayotganda) - eski parolni saqlab qolamiz
            conn.execute(
                "INSERT INTO hemis_accounts (user_id, domain, login, token, updated_at) "
                "VALUES (?, ?, ?, ?, ?) "
                "ON CONFLICT(user_id) DO UPDATE SET "
                "domain=excluded.domain, login=excluded.login, token=excluded.token, "
                "updated_at=excluded.updated_at",
                (user_id, domain, login, token, datetime.now().isoformat(timespec="seconds")),
            )
        conn.commit()


def db_get_hemis_account(user_id: int) -> dict | None:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute(
            "SELECT * FROM hemis_accounts WHERE user_id = ?", (user_id,)
        ).fetchone()
        return dict(row) if row else None


def db_delete_hemis_account(user_id: int) -> None:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        conn.execute("DELETE FROM hemis_accounts WHERE user_id = ?", (user_id,))
        conn.commit()


def db_update_hemis_token(user_id: int, token: str) -> None:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            "UPDATE hemis_accounts SET token = ?, updated_at = ? WHERE user_id = ?",
            (token, datetime.now().isoformat(timespec="seconds"), user_id),
        )
        conn.commit()


def db_add_admin(user_id: int, added_by: int | None = None) -> None:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            "INSERT OR IGNORE INTO admins (user_id, added_by, added_at) VALUES (?, ?, ?)",
            (user_id, added_by, datetime.now().isoformat(timespec="seconds")),
        )
        conn.commit()


def db_remove_admin(user_id: int) -> bool:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        cur = conn.execute("DELETE FROM admins WHERE user_id = ?", (user_id,))
        conn.commit()
        return cur.rowcount > 0


def db_get_dynamic_admin_ids() -> set[int]:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        rows = conn.execute("SELECT user_id FROM admins").fetchall()
        return {row[0] for row in rows}


def db_create_payment_request(
    user_id: int, username: str | None, full_name: str, service: str, amount: int, receipt_file_id: str
) -> int:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        cur = conn.execute(
            "INSERT INTO pending_payments "
            "(user_id, username, full_name, service, amount, receipt_file_id, status, created_at) "
            "VALUES (?, ?, ?, ?, ?, ?, 'pending', ?)",
            (user_id, username or "", full_name, service, amount, receipt_file_id,
             datetime.now().isoformat(timespec="seconds")),
        )
        conn.commit()
        return cur.lastrowid


def db_get_payment(payment_id: int) -> dict | None:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute(
            "SELECT * FROM pending_payments WHERE id = ?", (payment_id,)
        ).fetchone()
        return dict(row) if row else None


def db_set_payment_status(payment_id: int, status: str, decided_by: int) -> None:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            "UPDATE pending_payments SET status = ?, decided_at = ?, decided_by = ? WHERE id = ?",
            (status, datetime.now().isoformat(timespec="seconds"), decided_by, payment_id),
        )
        conn.commit()


def db_get_latest_payment_for_user(user_id: int) -> dict | None:
    with _db_lock, sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute(
            "SELECT * FROM pending_payments WHERE user_id = ? ORDER BY id DESC LIMIT 1",
            (user_id,),
        ).fetchone()
        return dict(row) if row else None


def is_admin(user_id: int) -> bool:
    return user_id in ADMIN_IDS or user_id in db_get_dynamic_admin_ids()


def _normalize_hemis_domain(text: str) -> str:
    """Foydalanuvchi kiritgan domenni to'liq bazaviy URL'ga aylantiradi.
    Masalan: 'student.tuit.uz' -> 'https://student.tuit.uz/rest/v1'
    Agar foydalanuvchi faqat universitet qisqartmasini yozsa (masalan 'TISU',
    'TERDU'), avtomatik 'student.<qisqartma>.uz' shakliga aylantiradi."""
    text = (text or "").strip().rstrip("/").lower()

    # Agar http/https va nuqta bo'lmasa - bu qisqartma (masalan "tisu")
    if not text.startswith("http://") and not text.startswith("https://") and "." not in text:
        text = f"student.{text}.uz"

    if not text.startswith("http://") and not text.startswith("https://"):
        text = "https://" + text
    if not text.endswith("/rest/v1"):
        text = text + "/rest/v1"
    return text


async def hemis_login(base_url: str, login: str, password: str) -> tuple[str | None, str]:
    """HEMIS'ga kirib, token oladi.
    Qaytaradi: (token yoki None, sabab_kodi).
    sabab_kodi: "ok" | "dns" | "auth" | "error"
    """
    try:
        async with aiohttp.ClientSession() as session:
            async with session.post(
                base_url + HEMIS_LOGIN_PATH,
                json={"login": login, "password": password},
                timeout=aiohttp.ClientTimeout(total=HEMIS_REQUEST_TIMEOUT),
            ) as resp:
                status = resp.status
                raw_text = await resp.text()
                if status != 200:
                    logging.warning(
                        "HEMIS login 200 emas. URL=%s status=%s javob=%s",
                        base_url + HEMIS_LOGIN_PATH, status, raw_text[:1000],
                    )
                    return None, "auth"
                try:
                    data = json.loads(raw_text)
                except Exception:
                    logging.warning(
                        "HEMIS login javobi JSON emas. URL=%s javob=%s",
                        base_url + HEMIS_LOGIN_PATH, raw_text[:1000],
                    )
                    return None, "auth"
    except aiohttp.ClientConnectorDNSError:
        logging.warning("HEMIS domeni topilmadi: %s", base_url)
        return None, "dns"
    except Exception:
        logging.exception("HEMIS'ga ulanishda xatolik")
        return None, "error"

    # Turli HEMIS instansiyalari tokenni har xil joyda qaytarishi mumkin
    inner = data.get("data") if isinstance(data, dict) else None
    token = (
        (inner or {}).get("token")
        if isinstance(inner, dict)
        else None
    ) or data.get("token") or data.get("access_token")
    if not token:
        logging.warning(
            "HEMIS login 200 qaytardi, lekin token topilmadi. URL=%s javob=%s",
            base_url + HEMIS_LOGIN_PATH, json.dumps(data)[:1000],
        )
        return None, "auth"
    return token, "ok"


def _format_hemis_login_error(reason: str) -> str:
    if reason == "dns":
        return (
            "❌ Bunday domen topilmadi. Universitet HEMIS manzilini tekshirib, "
            "qaytadan kiriting (masalan: <code>student.tesu.uz</code>)."
        )
    if reason == "auth":
        return "❌ Login yoki parol noto'g'ri. Qaytadan tekshirib kiriting."
    return "❌ HEMIS bilan bog'lanishda xatolik yuz berdi. Birozdan so'ng qaytadan urinib ko'ring."


async def hemis_api_get(base_url: str, token: str, path: str, params: dict | None = None):
    """HEMIS API'dan ma'lumot oladi. Qaytaradi: (data, http_status)."""
    headers = {"Authorization": f"Bearer {token}"}
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(
                base_url + path,
                headers=headers,
                params=params,
                timeout=aiohttp.ClientTimeout(total=HEMIS_REQUEST_TIMEOUT),
            ) as resp:
                status = resp.status
                try:
                    data = await resp.json(content_type=None)
                except Exception:
                    data = None
                return data, status
    except Exception:
        logging.exception("HEMIS API so'rovida xatolik")
        return None, 0


async def hemis_api_get_auto(user_id: int, account: dict, path: str, params: dict | None = None):
    """hemis_api_get bilan bir xil, lekin token muddati tugagan bo'lsa
    (401), saqlangan (shifrlangan) parol yordamida avtomatik qayta login
    qilib, qaytadan urinib ko'radi - foydalanuvchi hech narsa qayta
    yozmaydi."""
    base_url, token = account["domain"], account["token"]
    data, status = await hemis_api_get(base_url, token, path, params)
    if status != 401:
        return data, status

    password = _decrypt_text(account.get("password_enc"))
    if not password:
        return data, status  # parol saqlanmagan - qayta login imkoni yo'q

    new_token, _reason = await hemis_login(base_url, account["login"], password)
    if not new_token:
        return data, status

    db_update_hemis_token(user_id, new_token)
    return await hemis_api_get(base_url, new_token, path, params)


class UserTrackingMiddleware(BaseMiddleware):
    """Har bir yangilanishda foydalanuvchini SQLite bazasiga (bir marta) yozib boradi."""

    async def __call__(self, handler, event: TelegramObject, data: dict):
        user = data.get("event_from_user")
        if user is not None:
            try:
                db_add_user(user.id, user.username)
            except Exception:
                logging.exception("Foydalanuvchini bazaga yozishda xatolik")
        return await handler(event, data)


# ============================================================
# VALIDATSIYA YORDAMCHILARI
# ============================================================

def is_valid_phone(text: str) -> bool:
    """Telefon raqami faqat raqamlardan iborat va 9-13 xonali bo'lishi kerak."""
    text = (text or "").strip()
    return text.isdigit() and 9 <= len(text) <= 13


def is_valid_date(text: str) -> bool:
    """Sana 'kun.oy.yil' formatida (masalan 15.03.2000) ekanini tekshiradi.
    Oxirida qo'shimcha so'z bo'lishiga (masalan '15.03.2000-yil') ruxsat beriladi."""
    m = re.match(r"^\s*(\d{1,2})\.(\d{1,2})\.(\d{4})\s*(?:[-\s].*)?$", text or "")
    if not m:
        return False
    day, month, year = (int(x) for x in m.groups())
    try:
        datetime(year, month, day)
        return True
    except ValueError:
        return False


def _load_face_cascade() -> cv2.CascadeClassifier:
    """Haar cascade faylini bir nechta manbadan topishga harakat qiladi."""
    candidate_paths = []
    try:
        candidate_paths.append(cv2.data.haarcascades + "haarcascade_frontalface_default.xml")
    except AttributeError:
        pass

    cv2_dir = os.path.dirname(cv2.__file__)
    candidate_paths += [
        os.path.join(cv2_dir, "data", "haarcascade_frontalface_default.xml"),
        os.path.join(cv2_dir, "cv2", "data", "haarcascade_frontalface_default.xml"),
    ]

    for path in candidate_paths:
        if path and os.path.isfile(path):
            classifier = cv2.CascadeClassifier(path)
            if not classifier.empty():
                logging.info(f"Haar cascade topildi: {path}")
                return classifier

    raise RuntimeError(
        "Haar cascade fayli topilmadi yoki yuklanmadi.\n"
        "Yechim: pip uninstall opencv-python-headless -y && "
        "pip install opencv-python-headless==4.10.0.84"
    )


FACE_CASCADE = _load_face_cascade()

# rembg uchun yengil model ("u2netp", ~5MB) - standart "u2net" (~176MB) o'rniga.
# Xotira tejash uchun sessiya bir marta yaratiladi va qayta ishlatiladi
# (Railway kabi RAM-cheklangan hostinglarda process crash/OOM bo'lmasligi uchun muhim).
_REMBG_SESSION = new_session("u2netp")

def _build_bot() -> Bot:
    """PythonAnywhere bepul hisoblarida tashqi internetga faqat proxy orqali
    chiqish mumkin. Agar http(s)_proxy muhit o'zgaruvchisi mavjud bo'lsa (masalan
    PythonAnywhere avtomatik o'rnatgan bo'lsa), aiogram shu proxy orqali ishlaydi.
    Boshqa joyda (VPS, lokal kompyuter) bu o'zgaruvchi bo'lmaydi va oddiy
    to'g'ridan-to'g'ri ulanish ishlatiladi."""
    proxy_url = os.getenv("https_proxy") or os.getenv("HTTPS_PROXY") or os.getenv("http_proxy")
    if proxy_url:
        session = AiohttpSession(proxy=proxy_url)
        logging.info("Bot proxy orqali ulanmoqda: %s", proxy_url)
        return Bot(token=BOT_TOKEN, session=session)
    return Bot(token=BOT_TOKEN)


bot = _build_bot()
dp = Dispatcher(storage=MemoryStorage())
dp.update.outer_middleware(UserTrackingMiddleware())
db_init()  # Bazani polling yoki webhook rejimidan qat'iy nazar tayyorlab qo'yamiz


class FaceNotFoundError(Exception):
    """Rasmda yuz aniqlanmaganda ko'tariladi."""


def _is_low_quality(image_bytes: bytes) -> bool:
    """Juda kichik yoki loyqa rasmlarni oldindan aniqlab, rad etish uchun ishlatiladi."""
    try:
        img = Image.open(io.BytesIO(image_bytes))
        img = ImageOps.exif_transpose(img).convert("RGB")
    except Exception:
        return True

    w, h = img.size
    if w < MIN_IMAGE_DIMENSION or h < MIN_IMAGE_DIMENSION:
        return True

    gray = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2GRAY)
    variance = cv2.Laplacian(gray, cv2.CV_64F).var()
    return variance < BLUR_VARIANCE_THRESHOLD


def _center_crop_fallback(img: Image.Image, target_ratio: float) -> Image.Image:
    w, h = img.size
    current_ratio = w / h
    if current_ratio > target_ratio:
        new_w = int(h * target_ratio)
        left = (w - new_w) // 2
        return img.crop((left, 0, left + new_w, h))
    else:
        new_h = int(w / target_ratio)
        top = max(0, (h - new_h) // 3)
        return img.crop((0, top, w, top + new_h))


def _detect_face(img: Image.Image):
    """Rasmdagi eng katta yuzni qaytaradi: (fx, fy, fw, fh) yoki None."""
    cv_img = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR)
    gray = cv2.cvtColor(cv_img, cv2.COLOR_BGR2GRAY)
    faces = FACE_CASCADE.detectMultiScale(
        gray, scaleFactor=1.1, minNeighbors=5, minSize=(60, 60)
    )
    if len(faces) == 0:
        return None
    return max(faces, key=lambda f: f[2] * f[3])


def make_3x4(image_bytes: bytes) -> Image.Image:
    """Yuzni aniqlab, hujjatbop 3x4 kompozitsiyaga moslab kesadi. PIL Image qaytaradi."""
    img = Image.open(io.BytesIO(image_bytes))
    img = ImageOps.exif_transpose(img)
    img = img.convert("RGB")

    target_ratio = TARGET_WIDTH / TARGET_HEIGHT
    w, h = img.size

    face = _detect_face(img)

    if face is None:
        raise FaceNotFoundError("Rasmda yuz aniqlanmadi")
    else:
        fx, fy, fw, fh = face
        head_top = fy - 0.35 * fh
        head_bottom = fy + fh + 0.25 * fh
        head_height = head_bottom - head_top
        face_cx = fx + fw / 2

        crop_h = head_height / HEAD_HEIGHT_RATIO
        crop_w = crop_h * target_ratio

        top_margin = crop_h * 0.14
        crop_top = head_top - top_margin
        crop_bottom = crop_top + crop_h
        crop_left = face_cx - crop_w / 2
        crop_right = crop_left + crop_w

        if crop_left < 0:
            crop_right -= crop_left
            crop_left = 0
        if crop_right > w:
            shift = crop_right - w
            crop_left = max(0, crop_left - shift)
            crop_right = w
        if crop_top < 0:
            crop_bottom -= crop_top
            crop_top = 0
        if crop_bottom > h:
            shift = crop_bottom - h
            crop_top = max(0, crop_top - shift)
            crop_bottom = h

        crop_left, crop_top = int(crop_left), int(crop_top)
        crop_right, crop_bottom = int(crop_right), int(crop_bottom)
        cropped = img.crop((crop_left, crop_top, crop_right, crop_bottom))

        cur_ratio = cropped.width / cropped.height
        if abs(cur_ratio - target_ratio) > 0.01:
            cropped = _center_crop_fallback(cropped, target_ratio)

    return cropped.resize((TARGET_WIDTH, TARGET_HEIGHT), Image.LANCZOS)


def remove_bg_to_white(img: Image.Image) -> Image.Image:
    """Haqiqiy fonni olib tashlab, oq fon bilan almashtiradi."""
    # Xotira sarfini kamaytirish uchun: rasm baribir 3x4ga kichraytiriladi,
    # shuning uchun rembg'ga yuborishdan oldin ortiqcha katta rasmni
    # cheklab qo'yamiz (RAM-cheklangan hostinglarda OOM bo'lmasligi uchun).
    MAX_REMBG_DIMENSION = 1200
    if max(img.size) > MAX_REMBG_DIMENSION:
        img = img.copy()
        img.thumbnail((MAX_REMBG_DIMENSION, MAX_REMBG_DIMENSION), Image.LANCZOS)

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    result_bytes = remove(buf.getvalue(), session=_REMBG_SESSION)  # shaffof fonli RGBA PNG

    fg = Image.open(io.BytesIO(result_bytes)).convert("RGBA")

    # Chekka atrofida qolib ketadigan "halo/nur" effektini yo'qotish:
    # shaffoflikni qattiq chegaraga (0 yoki 255) o'tkazamiz, so'ng bir oz
    # eroziya qilib qolgan nozik chekkani ham kesib tashlaymiz.
    r, g, b, a = fg.split()
    a = a.point(lambda x: 255 if x > 140 else 0)
    a = a.filter(ImageFilter.MinFilter(7))
    fg = Image.merge("RGBA", (r, g, b, a))

    white_bg = Image.new("RGBA", fg.size, (255, 255, 255, 255))
    white_bg.alpha_composite(fg)
    return white_bg.convert("RGB")


def _to_jpeg_bytes(img: Image.Image) -> bytes:
    out = io.BytesIO()
    img.save(out, format="JPEG", quality=95, dpi=(300, 300))
    out.seek(0)
    return out.read()


# Vaqtinchalik xotira: user_id -> oxirgi tayyor 3x4 rasm (chop etish
# varag'i yaratish uchun).
pending_photos: dict[int, Image.Image] = {}

# Vaqtinchalik xotira: user_id -> to'lov tasdiqlanishini kutayotgan XOM rasm
# baytlari (3x4 uchun). Admin tasdiqlagach shu bayt asosida qayta ishlanadi.
pending_3x4_payload: dict[int, bytes] = {}

# Vaqtinchalik xotira: user_id -> to'lov tasdiqlanishini kutayotgan
# ma'lumotnoma ma'lumotlari va rasmi. Admin tasdiqlagach hujjat shundan tuziladi.
pending_obyektivka_payload: dict[int, dict] = {}


def _print_sheet_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="🖨 Chop etish varag'i (8 dona)", callback_data="print_sheet")],
        ]
    )


def make_print_sheet(photo: Image.Image, cols: int = 4, rows: int = 2,
                      gap: int = 40, margin: int = 30,
                      border_width: int = 2) -> Image.Image:
    """Bitta 3x4 rasmdan ko'p nusxali chop etish varag'ini yasaydi
    (fotostudiyalarda ishlatiladigan "8 dona 3x4" varag'iga o'xshab).
    Har bir nusxa atrofida qora ramka chiziladi."""
    pw, ph = photo.size
    canvas_w = margin * 2 + pw * cols + gap * (cols - 1)
    canvas_h = margin * 2 + ph * rows + gap * (rows - 1)
    canvas = Image.new("RGB", (canvas_w, canvas_h), (255, 255, 255))
    draw = ImageDraw.Draw(canvas)

    for r in range(rows):
        for c in range(cols):
            x = margin + c * (pw + gap)
            y = margin + r * (ph + gap)
            canvas.paste(photo, (x, y))
            draw.rectangle(
                [x - border_width, y - border_width, x + pw + border_width - 1, y + ph + border_width - 1],
                outline=(0, 0, 0),
                width=border_width,
            )
    return canvas


def frame_photo(photo: Image.Image, border_width: int = 5, color=(0, 0, 0)) -> Image.Image:
    """3x4 rasm atrofiga qora ramka chizadi (rasm o'lchamini o'zgartirmaydi,
    faqat chetlariga chiziq tortadi) — ma'lumotnomaga qo'yiladigan rasm
    hujjatbop ramkali ko'rinishda bo'lishi uchun."""
    framed = photo.copy()
    draw = ImageDraw.Draw(framed)
    w, h = framed.size
    draw.rectangle([0, 0, w - 1, h - 1], outline=color, width=border_width)
    return framed


# ============================================================
# ASOSIY MENYU (Reply keyboard) VA ADMIN PANEL
# ============================================================

def main_menu_kb(user_id: int) -> ReplyKeyboardMarkup:
    rows = [
        [KeyboardButton(text="📸 3x4 rasm"), KeyboardButton(text="🗂 Ma'lumotnoma")],
        [KeyboardButton(text="📋 Namuna"), KeyboardButton(text="📊 Holatim")],
        [KeyboardButton(text="🎓 HEMIS"), KeyboardButton(text="❓ Yordam")],
        [KeyboardButton(text="❌ Bekor qilish")],
    ]
    if is_admin(user_id):
        rows.append([KeyboardButton(text="⚙️ Admin panel")])
    return ReplyKeyboardMarkup(keyboard=rows, resize_keyboard=True)


def main_inline_menu_kb(user_id: int) -> InlineKeyboardMarkup:
    rows = [
        [
            InlineKeyboardButton(text="📸 3x4 rasm", callback_data="nav_3x4"),
            InlineKeyboardButton(text="🗂 Ma'lumotnoma", callback_data="nav_malumotnoma"),
        ],
        [
            InlineKeyboardButton(text="📋 Namuna", callback_data="nav_namuna"),
            InlineKeyboardButton(text="📊 Holatim", callback_data="nav_holatim"),
        ],
        [
            InlineKeyboardButton(text="🎓 HEMIS", callback_data="nav_hemis_menu"),
            InlineKeyboardButton(text="❓ Yordam", callback_data="nav_yordam"),
        ],
    ]
    if is_admin(user_id):
        rows.append([InlineKeyboardButton(text="⚙️ Admin panel", callback_data="nav_admin")])
    return InlineKeyboardMarkup(inline_keyboard=rows)



def hemis_submenu_kb(user_id: int) -> InlineKeyboardMarkup:
    account = db_get_hemis_account(user_id)
    rows = []
    if account:
        rows.append([InlineKeyboardButton(text="📊 Baholarim", callback_data="nav_baholarim")])
        rows.append([InlineKeyboardButton(text="🔌 Hisobni uzish", callback_data="nav_hemis_uzish")])
    else:
        rows.append([InlineKeyboardButton(text="🔗 Hisobni ulash", callback_data="nav_hemis_connect")])
    rows.append([InlineKeyboardButton(text="🔙 Bosh menyu", callback_data="nav_main")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def _back_to_menu_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[[InlineKeyboardButton(text="🔙 Bosh menyu", callback_data="nav_main")]]
    )


async def _show_menu_section(call: CallbackQuery, text: str, kb: InlineKeyboardMarkup) -> None:
    """Joriy xabarni tahrirlashga urinadi; agar bo'lmasa (masalan xabar hujjat/rasm
    bo'lsa), yangi xabar yuboradi."""
    try:
        await call.message.edit_text(text, reply_markup=kb, parse_mode="HTML")
    except Exception:
        await call.message.answer(text, reply_markup=kb, parse_mode="HTML")


def admin_panel_kb(back: bool = False) -> InlineKeyboardMarkup:
    rows = [
        [InlineKeyboardButton(text="📊 Statistika", callback_data="admin_stats")],
        [InlineKeyboardButton(text="📢 Xabar yuborish", callback_data="admin_broadcast_cb")],
        [InlineKeyboardButton(text="👥 Adminlar ro'yxati", callback_data="admin_list_cb")],
        [InlineKeyboardButton(text="🛠 Admin/Balans buyruqlari", callback_data="admin_help_cb")],
    ]
    if back:
        rows.append([InlineKeyboardButton(text="🔙 Bosh menyu", callback_data="nav_main")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def _format_stats(stats: dict, pay_stats: dict) -> str:
    by_service = pay_stats["by_service"]
    lines = [
        "📊 <b>Bot statistikasi</b>\n",
        f"👥 Jami foydalanuvchilar: <b>{stats['total_users']}</b>",
        f"📄 Jami ma'lumotnoma so'rovlari: <b>{stats['total_requests']}</b>",
        f"📅 Bugungi so'rovlar: <b>{stats['today_requests']}</b>",
        f"⚠️ Xatoliklar soni: <b>{stats['total_errors']}</b>",
        "",
        "💳 <b>To'lovlar</b>",
        f"⏳ Kutilmoqda: <b>{pay_stats['pending']}</b>",
        f"✅ Tasdiqlangan: <b>{pay_stats['confirmed']}</b>",
        f"❌ Rad etilgan: <b>{pay_stats['rejected']}</b>",
        f"   • 📸 3x4: <b>{by_service.get('3x4', 0)}</b>",
        f"   • 🗂 Ma'lumotnoma: <b>{by_service.get('obyektivka', 0)}</b>",
        "",
        f"💰 Jami tushum: <b>{pay_stats['revenue']:,} so'm</b>".replace(",", " "),
        f"💰 Bugungi tushum: <b>{pay_stats['revenue_today']:,} so'm</b>".replace(",", " "),
    ]
    return "\n".join(lines)


@dp.message(CommandStart())
async def start_handler(message: Message):
    await message.answer(
        "Salom! 📸\n\n"
        "Menga oddiy rasm yuboring — men uni 3x4 sm (300 DPI) hujjatbop "
        "o'lchamga moslab, fonini oqartirib beraman.\n\n"
        "Rasmiy ma'lumotnoma (obyektivka) tuzish uchun /malumotnoma buyrug'ini yuboring.\n"
        "Ma'lumotnoma qanday ko'rinishda bo'lishini oldindan ko'rish uchun /namuna.\n"
        "Jarayonni istalgan vaqtda /bekor bilan to'xtatishingiz mumkin.\n\n"
        "Eng yaxshi natija uchun: yuz aniq ko'rinadigan rasm yuboring.\n\n"
        f"💳 <b>Narxlar:</b>\n"
        f"📸 3x4 rasm — {PRICE_3X4:,} so'm\n"
        f"🗂 Ma'lumotnoma — {PRICE_OBYEKTIVKA:,} so'm\n"
        "Rasm/ma'lumotlaringizni yuborganingizdan so'ng, avval to'lov qilib "
        "chekini botga tashlashingiz kerak bo'ladi. Admin tekshirib tasdiqlagach, "
        "tayyor natija sizga avtomatik yuboriladi.".replace(",", " "),
        reply_markup=main_menu_kb(message.from_user.id),
        parse_mode="HTML",
    )


@dp.message(F.text == "📸 3x4 rasm")
async def menu_3x4_photo(message: Message):
    await message.answer(
        "📸 3x4 rasm olish uchun menga oddiy fotosuratingizni (yuzingiz aniq "
        "ko'rinadigan) rasm sifatida yuboring — men uni 3x4 sm (300 DPI) hujjatbop "
        "o'lchamga moslab, fonini oqartirib beraman."
    )


HELP_TEXT = (
    "❓ <b>Yordam</b>\n\n"
    "☰ /menu — bosh menyu (tugmalar orqali)\n"
    "📸 /start — botni ishga tushirish\n"
    "🗂 /malumotnoma — ma'lumotnoma (obyektivka) tuzish\n"
    "📋 /namuna — ma'lumotnoma namunasini ko'rish\n"
    "📊 /holatim — oxirgi to'lov so'rovingiz holati\n"
    "🎓 /hemis — HEMIS hisobini ulash\n"
    "📈 /baholarim — HEMIS orqali baho/GPA ko'rish\n"
    "🔌 /hemis_uzish — HEMIS hisobini uzish\n"
    "❌ /bekor — joriy jarayonni bekor qilish\n\n"
    "💳 3x4 rasm yoki ma'lumotnoma uchun avval to'lov qilib chek yuborishingiz "
    "kerak. Admin tasdiqlagach, natija avtomatik yuboriladi."
)


def _format_holatim_text(user_id: int) -> str:
    payment = db_get_latest_payment_for_user(user_id)
    if not payment:
        return "ℹ️ Sizda hali birorta ham to'lov so'rovi yo'q."

    status_map = {
        "pending": "⏳ Kutilmoqda",
        "confirmed": "✅ Tasdiqlangan",
        "rejected": "❌ Rad etilgan",
    }
    name = SERVICE_NAMES.get(payment["service"], payment["service"])
    status_text = status_map.get(payment["status"], payment["status"])
    lines = [
        "📊 <b>Oxirgi so'rovingiz holati</b>\n",
        f"🛍 Xizmat: {name}",
        f"💵 Summasi: {payment['amount']:,} so'm".replace(",", " "),
        f"📌 Holati: <b>{status_text}</b>",
        f"🕒 Yuborilgan: {payment['created_at']}",
    ]
    if payment["decided_at"]:
        lines.append(f"✅ Ko'rib chiqilgan: {payment['decided_at']}")
    return "\n".join(lines)


@dp.message(Command("yordam"))
async def yordam_command(message: Message):
    await message.answer(HELP_TEXT, parse_mode="HTML")


@dp.message(Command("holatim"))
async def holatim_command(message: Message):
    await message.answer(_format_holatim_text(message.from_user.id), parse_mode="HTML")


@dp.message(Command("menu"))
async def menu_command(message: Message):
    await message.answer(
        "📋 <b>Bosh menyu</b>\n\nKerakli bo'limni tanlang:",
        reply_markup=main_inline_menu_kb(message.from_user.id),
        parse_mode="HTML",
    )


@dp.message(F.text == "☰ Bosh menyu")
async def menu_button(message: Message):
    await menu_command(message)


@dp.callback_query(F.data == "nav_main")
async def nav_main_callback(call: CallbackQuery):
    await _show_menu_section(
        call,
        "📋 <b>Bosh menyu</b>\n\nKerakli bo'limni tanlang:",
        main_inline_menu_kb(call.from_user.id),
    )
    await call.answer()


@dp.callback_query(F.data == "nav_3x4")
async def nav_3x4_callback(call: CallbackQuery):
    text = (
        "📸 3x4 rasm olish uchun menga oddiy fotosuratingizni (yuzingiz aniq "
        "ko'rinadigan) rasm sifatida yuboring — men uni 3x4 sm (300 DPI) hujjatbop "
        "o'lchamga moslab, fonini oqartirib beraman."
    )
    await _show_menu_section(call, text, _back_to_menu_kb())
    await call.answer()


@dp.callback_query(F.data == "nav_namuna")
async def nav_namuna_callback(call: CallbackQuery):
    await call.answer()
    await _send_sample_document(call.message)
    await call.message.answer("⬆️ Namuna yuqorida.", reply_markup=_back_to_menu_kb())


@dp.callback_query(F.data == "nav_malumotnoma")
async def nav_malumotnoma_callback(call: CallbackQuery, state: FSMContext):
    await call.answer()
    await _start_malumotnoma_flow(call.message, call.from_user.id, state)


@dp.callback_query(F.data == "nav_holatim")
async def nav_holatim_callback(call: CallbackQuery):
    text = _format_holatim_text(call.from_user.id)
    await _show_menu_section(call, text, _back_to_menu_kb())
    await call.answer()


@dp.callback_query(F.data == "nav_yordam")
async def nav_yordam_callback(call: CallbackQuery):
    await _show_menu_section(call, HELP_TEXT, _back_to_menu_kb())
    await call.answer()


@dp.callback_query(F.data == "nav_admin")
async def nav_admin_callback(call: CallbackQuery):
    if not is_admin(call.from_user.id):
        await call.answer("⛔️ Ruxsat yo'q.", show_alert=True)
        return
    await _show_menu_section(call, "⚙️ <b>Admin panel</b>", admin_panel_kb(back=True))
    await call.answer()


@dp.callback_query(F.data == "nav_hemis_menu")
async def nav_hemis_menu_callback(call: CallbackQuery):
    if not HEMIS_SERVICE_ENABLED:
        await _show_menu_section(call, HEMIS_UNAVAILABLE_TEXT, _back_to_menu_kb())
        await call.answer()
        return
    account = db_get_hemis_account(call.from_user.id)
    if account:
        text = (
            f"🎓 <b>HEMIS</b>\n\nSiz <code>{account['domain']}</code> ga "
            f"<code>{account['login']}</code> hisobi bilan ulangansiz."
        )
    else:
        text = "🎓 <b>HEMIS</b>\n\nHisobingizni ulab, baho va GPA ma'lumotlaringizni ko'rishingiz mumkin."
    await _show_menu_section(call, text, hemis_submenu_kb(call.from_user.id))
    await call.answer()


@dp.callback_query(F.data == "nav_hemis_connect")
async def nav_hemis_connect_callback(call: CallbackQuery, state: FSMContext):
    await call.answer()
    await _hemis_connect_flow(call.message, call.from_user.id, state)


@dp.callback_query(F.data == "nav_baholarim")
async def nav_baholarim_callback(call: CallbackQuery):
    await call.answer()
    await _send_baholarim(call.message, call.from_user.id)


@dp.callback_query(F.data == "nav_hemis_uzish")
async def nav_hemis_uzish_callback(call: CallbackQuery):
    db_delete_hemis_account(call.from_user.id)
    await _show_menu_section(
        call, "✅ HEMIS hisobingiz uzildi.", hemis_submenu_kb(call.from_user.id)
    )
    await call.answer()


@dp.message(Command("admin"))
async def admin_command(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("⛔️ Bu buyruq faqat admin uchun.")
        return
    await message.answer("⚙️ <b>Admin panel</b>", reply_markup=admin_panel_kb(), parse_mode="HTML")


@dp.message(F.text == "⚙️ Admin panel")
async def admin_menu_button(message: Message):
    if not is_admin(message.from_user.id):
        return
    await message.answer("⚙️ <b>Admin panel</b>", reply_markup=admin_panel_kb(), parse_mode="HTML")


@dp.message(Command("statistika"))
async def statistika_command(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("⛔️ Bu buyruq faqat admin uchun.")
        return
    stats = db_get_stats()
    pay_stats = db_get_payment_stats()
    await message.answer(_format_stats(stats, pay_stats), parse_mode="HTML")


@dp.callback_query(F.data == "admin_stats")
async def admin_stats_callback(call: CallbackQuery):
    if not is_admin(call.from_user.id):
        await call.answer("⛔️ Ruxsat yo'q.", show_alert=True)
        return
    stats = db_get_stats()
    pay_stats = db_get_payment_stats()
    await call.message.answer(_format_stats(stats, pay_stats), parse_mode="HTML")
    await call.answer()


@dp.callback_query(F.data == "admin_broadcast_cb")
async def admin_broadcast_callback(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id):
        await call.answer("⛔️ Ruxsat yo'q.", show_alert=True)
        return
    await state.set_state(Broadcast.waiting_content)
    await call.message.answer(
        "📢 Barcha foydalanuvchilarga yuboriladigan xabarni yuboring "
        "(matn, rasm yoki hujjat bo'lishi mumkin).\n"
        "Bekor qilish uchun /bekor buyrug'ini yuboring."
    )
    await call.answer()


@dp.callback_query(F.data == "admin_list_cb")
async def admin_list_callback(call: CallbackQuery):
    if not is_admin(call.from_user.id):
        await call.answer("⛔️ Ruxsat yo'q.", show_alert=True)
        return
    lines = ["👥 <b>Adminlar ro'yxati</b>\n"]
    for admin_id in sorted(ADMIN_IDS):
        lines.append(f"• <code>{admin_id}</code> (asosiy admin)")
    for admin_id in sorted(db_get_dynamic_admin_ids()):
        lines.append(f"• <code>{admin_id}</code> (qo'shilgan admin)")
    await call.message.answer("\n".join(lines), parse_mode="HTML")
    await call.answer()


@dp.callback_query(F.data == "admin_help_cb")
async def admin_help_callback(call: CallbackQuery):
    if not is_admin(call.from_user.id):
        await call.answer("⛔️ Ruxsat yo'q.", show_alert=True)
        return
    await call.message.answer(
        "🛠 <b>Admin buyruqlari</b>\n\n"
        "➕ /admin_qoshish &lt;user_id&gt; — yangi admin qo'shish\n"
        "➖ /admin_ochirish &lt;user_id&gt; — adminlikdan olib tashlash\n"
        "👥 /adminlar — barcha adminlar ro'yxati\n"
        "📢 /xabar — barcha foydalanuvchilarga xabar yuborish (matn/rasm/hujjat)\n\n"
        "💳 <b>To'lovlar qanday ishlaydi:</b>\n"
        "Foydalanuvchi 3x4 rasm yuboradi yoki ma'lumotnoma so'rovnomasini "
        "to'ldiradi — natija darhol berilmaydi, avval unga karta raqami "
        "ko'rsatiladi va chek yuborish tugmasi chiqadi. Chek yuborilganda, "
        "u sizga (barcha adminlarga) rasm va \"✅ Tasdiqlash\" / \"❌ Rad etish\" "
        "tugmalari bilan yuboriladi. Tasdiqlasangiz — tayyor rasm/hujjat "
        "foydalanuvchiga avtomatik yuboriladi; rad etsangiz — hech narsa "
        "yuborilmaydi va u qaytadan urinishi kerak bo'ladi.\n\n"
        f"📌 Joriy narxlar: 3x4 rasm — {PRICE_3X4:,} so'm, "
        f"Ma'lumotnoma — {PRICE_OBYEKTIVKA:,} so'm".replace(",", " "),
        parse_mode="HTML",
    )
    await call.answer()


@dp.message(Command("admin_qoshish"))
async def admin_qoshish_handler(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("⛔️ Bu buyruq faqat admin uchun.")
        return
    parts = (message.text or "").split()
    if len(parts) < 2 or not parts[1].lstrip("-").isdigit():
        await message.answer(
            "Foydalanish: /admin_qoshish <user_id>\n"
            "Masalan: /admin_qoshish 123456789\n\n"
            "user_id ni bilish uchun foydalanuvchi @userinfobot ga yozishi mumkin."
        )
        return
    new_admin_id = int(parts[1])
    if is_admin(new_admin_id):
        await message.answer("ℹ️ Bu foydalanuvchi allaqachon admin.")
        return
    db_add_admin(new_admin_id, added_by=message.from_user.id)
    await _set_admin_commands_for(new_admin_id)
    await message.answer(f"✅ {new_admin_id} endi admin sifatida qo'shildi (barcha xizmatlar bepul).")
    try:
        await bot.send_message(
            new_admin_id,
            "🎉 Sizga admin huquqi berildi! Endi botning barcha xizmatlaridan "
            "(3x4 rasm, ma'lumotnoma) bepul foydalanishingiz mumkin.\n"
            "/admin buyrug'i orqali admin panelga kirishingiz mumkin.",
        )
    except Exception:
        logging.info("Yangi adminga xabar yuborib bo'lmadi (bot bilan hali /start bosmagan bo'lishi mumkin).")


@dp.message(Command("admin_ochirish"))
async def admin_ochirish_handler(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("⛔️ Bu buyruq faqat admin uchun.")
        return
    parts = (message.text or "").split()
    if len(parts) < 2 or not parts[1].lstrip("-").isdigit():
        await message.answer("Foydalanish: /admin_ochirish <user_id>")
        return
    target_id = int(parts[1])
    if target_id in ADMIN_IDS:
        await message.answer(
            "⚠️ Bu admin dasturiy kod ichida (ADMIN_IDS) belgilangan, "
            "uni shu buyruq bilan olib tashlab bo'lmaydi."
        )
        return
    removed = db_remove_admin(target_id)
    if removed:
        try:
            await bot.set_my_commands(DEFAULT_COMMANDS, scope=BotCommandScopeChat(chat_id=target_id))
        except Exception:
            pass
        await message.answer(f"✅ {target_id} adminlikdan olib tashlandi.")
    else:
        await message.answer("ℹ️ Bu foydalanuvchi qo'shimcha adminlar ro'yxatida topilmadi.")


@dp.message(Command("adminlar"))
async def adminlar_handler(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("⛔️ Bu buyruq faqat admin uchun.")
        return
    lines = ["👥 <b>Adminlar ro'yxati</b>\n"]
    for admin_id in sorted(ADMIN_IDS):
        lines.append(f"• <code>{admin_id}</code> (asosiy admin)")
    for admin_id in sorted(db_get_dynamic_admin_ids()):
        lines.append(f"• <code>{admin_id}</code> (qo'shilgan admin)")
    await message.answer("\n".join(lines), parse_mode="HTML")


class Broadcast(StatesGroup):
    waiting_content = State()
    waiting_confirm = State()


class HemisAuth(StatesGroup):
    waiting_domain = State()
    waiting_login = State()
    waiting_password = State()


def _hemis_connect_kb() -> ReplyKeyboardMarkup | None:
    """Agar mini-app manzili (HEMIS_WEBAPP_URL) sozlangan bo'lsa, forma
    ochadigan tugma qaytaradi. Aks holda None - bu holda chat orqali
    (matn yozib) ulanish oqimi ishlatiladi."""
    if HEMIS_WEBAPP_URL:
        return ReplyKeyboardMarkup(
            keyboard=[
                [KeyboardButton(text="🔐 HEMIS'ga kirish", web_app=WebAppInfo(url=HEMIS_WEBAPP_URL))]
            ],
            resize_keyboard=True,
            one_time_keyboard=True,
        )
    return None


async def _hemis_connect_flow(send_target, user_id: int, state: FSMContext) -> None:
    if not HEMIS_SERVICE_ENABLED:
        await send_target.answer(HEMIS_UNAVAILABLE_TEXT)
        return
    account = db_get_hemis_account(user_id)
    if account:
        await send_target.answer(
            f"ℹ️ Siz allaqachon <code>{account['domain']}</code> ga "
            f"<code>{account['login']}</code> hisobi bilan ulangansiz.\n"
            "Qayta ulanish uchun avval /hemis_uzish buyrug'ini yuboring.",
            parse_mode="HTML",
        )
        return

    kb = _hemis_connect_kb()
    if kb:
        await send_target.answer(
            "🎓 HEMIS hisobingizni ulash uchun quyidagi tugmani bosing va "
            "ochilgan oynada universitet domeningiz, login (talaba ID) va "
            "parolingizni kiriting. Ma'lumotlaringiz xavfsiz, shifrlangan "
            "holda saqlanadi va keyingi safar qayta kiritish shart emas.",
            reply_markup=kb,
        )
        return

    # Mini-app manzili hali sozlanmagan (masalan lokal/polling test rejimi) -
    # shu holatda ma'lumotlarni chatga yozib kiritamiz.
    await state.set_state(HemisAuth.waiting_domain)
    await send_target.answer(
        "🎓 HEMIS hisobingizni ulaymiz.\n\n"
        "Universitetingizning HEMIS talaba portali domenini yuboring "
        "(masalan: <code>student.tuit.uz</code>).\n"
        "Bekor qilish uchun /bekor.",
        parse_mode="HTML",
    )


@dp.message(Command("hemis"))
async def hemis_command(message: Message, state: FSMContext):
    await _hemis_connect_flow(message, message.from_user.id, state)


@dp.message(F.web_app_data)
async def hemis_webapp_data_handler(message: Message):
    """Mini-app (WebApp) formasidan yuborilgan login/parolni qabul qiladi."""
    logging.info(f"WEBAPP DATA KELDI: user={message.from_user.id}, raw={message.web_app_data.data!r}")
    if not HEMIS_SERVICE_ENABLED:
        await message.answer(HEMIS_UNAVAILABLE_TEXT)
        return
    try:
        payload = json.loads(message.web_app_data.data)
    except Exception:
        await message.answer("⚠️ Ma'lumotni o'qib bo'lmadi, qaytadan urinib ko'ring.")
        return

    if payload.get("type") != "hemis_login":
        return

    domain_text = (payload.get("domain") or "").strip()
    login = (payload.get("login") or "").strip()
    password = payload.get("password") or ""

    if not domain_text or not login or not password:
        await message.answer("⚠️ Barcha maydonlarni to'ldiring va qaytadan urinib ko'ring.")
        return

    base_url = _normalize_hemis_domain(domain_text)
    wait_msg = await message.answer("⏳ HEMIS tizimiga ulanmoqda...")
    token, reason = await hemis_login(base_url, login, password)

    if not token:
        await wait_msg.edit_text(
            _format_hemis_login_error(reason) + "\n\n/hemis orqali qaytadan urinib ko'ring.",
            parse_mode="HTML",
        )
        return

    db_set_hemis_account(message.from_user.id, base_url, login, token, password=password)
    await wait_msg.edit_text(
        "✅ HEMIS hisobingiz muvaffaqiyatli ulandi va xavfsiz saqlandi!\n"
        "Endi /baholarim orqali, qayta login qilmasdan ma'lumotlaringizni ko'rishingiz mumkin."
    )


@dp.message(Command("hemis_uzish"))
async def hemis_disconnect_command(message: Message):
    if not HEMIS_SERVICE_ENABLED:
        await message.answer(HEMIS_UNAVAILABLE_TEXT)
        return
    db_delete_hemis_account(message.from_user.id)
    await message.answer("✅ HEMIS hisobingiz uzildi.")


@dp.message(HemisAuth.waiting_domain)
async def hemis_domain_handler(message: Message, state: FSMContext):
    domain_text = (message.text or "").strip()
    if not domain_text or " " in domain_text:
        await message.answer("Iltimos, faqat domen nomini yuboring (masalan: student.tuit.uz).")
        return
    base_url = _normalize_hemis_domain(domain_text)
    await state.update_data(_hemis_base_url=base_url)
    await state.set_state(HemisAuth.waiting_login)
    await message.answer("👤 Endi HEMIS login (talaba ID)ingizni yuboring:")


@dp.message(HemisAuth.waiting_login)
async def hemis_login_handler(message: Message, state: FSMContext):
    login_text = (message.text or "").strip()
    if not login_text:
        await message.answer("Iltimos, login (talaba ID)ingizni matn ko'rinishida yuboring.")
        return
    await state.update_data(_hemis_login=login_text)
    await state.set_state(HemisAuth.waiting_password)
    await message.answer(
        "🔒 Endi HEMIS parolingizni yuboring.\n"
        "Xabar yuborilgandan so'ng darhol chatdan o'chiriladi; parolingiz "
        "shifrlangan holda saqlanadi (keyin qayta yozmasligingiz uchun)."
    )


@dp.message(HemisAuth.waiting_password)
async def hemis_password_handler(message: Message, state: FSMContext):
    password = (message.text or "").strip()
    data = await state.get_data()
    base_url = data.get("_hemis_base_url")
    login = data.get("_hemis_login")

    # Parolni chatdan darhol o'chirib tashlaymiz
    try:
        await message.delete()
    except Exception:
        pass

    if not password or not base_url or not login:
        await state.clear()
        await message.answer("⚠️ Xatolik yuz berdi, /hemis orqali qaytadan urinib ko'ring.")
        return

    wait_msg = await message.answer("⏳ HEMIS tizimiga ulanmoqda...")
    token, reason = await hemis_login(base_url, login, password)
    await state.clear()

    if not token:
        await wait_msg.edit_text(
            _format_hemis_login_error(reason)
            + "\n\nQaytadan urinish uchun /hemis buyrug'ini yuboring.",
            parse_mode="HTML",
        )
        return

    db_set_hemis_account(message.from_user.id, base_url, login, token, password=password)
    await wait_msg.edit_text(
        "✅ HEMIS hisobingiz muvaffaqiyatli ulandi va xavfsiz saqlandi!\n"
        "Endi /baholarim buyrug'i orqali baho va GPA ma'lumotlaringizni ko'rishingiz mumkin."
    )




def _format_hemis_error(status: int) -> str:
    if status == 401:
        return (
            "🔒 Sessiya muddati tugagan. Iltimos, /hemis_uzish, so'ng /hemis "
            "orqali qaytadan ulaning."
        )
    return "⚠️ HEMIS'dan ma'lumot olishda xatolik yuz berdi. Birozdan so'ng qaytadan urinib ko'ring."


async def _send_baholarim(send_target, user_id: int) -> None:
    if not HEMIS_SERVICE_ENABLED:
        await send_target.answer(HEMIS_UNAVAILABLE_TEXT)
        return
    account = db_get_hemis_account(user_id)
    if not account:
        await send_target.answer(
            "Avval HEMIS hisobingizni ulang.",
            reply_markup=_hemis_connect_kb(),
        )
        return

    gpa_data, gpa_status = await hemis_api_get_auto(user_id, account, HEMIS_GPA_PATH)
    if gpa_status == 401 or gpa_status == 0 or gpa_data is None:
        await send_target.answer(
            _format_hemis_error(gpa_status), reply_markup=_hemis_connect_kb()
        )
        return

    gpa_list = (gpa_data or {}).get("data", []) if isinstance(gpa_data, dict) else []
    lines = ["🎓 <b>GPA ma'lumotlaringiz</b>\n"]
    if gpa_list:
        for item in gpa_list:
            level = item.get("level", {}).get("name") if isinstance(item.get("level"), dict) else item.get("level", "")
            gpa_val = item.get("gpa", item.get("avg_gpa", "—"))
            lines.append(f"• {level}: <b>{gpa_val}</b>")
    else:
        lines.append("Ma'lumot topilmadi.")

    subjects_data, sub_status = await hemis_api_get_auto(user_id, account, HEMIS_SUBJECTS_PATH)
    if sub_status == 200 and isinstance(subjects_data, dict):
        subjects = subjects_data.get("data", [])
        if subjects:
            lines.append("\n📚 <b>Joriy semestr fanlari</b>\n")
            for s in subjects[:20]:
                name = s.get("subject", {}).get("name") if isinstance(s.get("subject"), dict) else s.get("subject", "")
                grade = s.get("grade", s.get("total_ball", "—"))
                lines.append(f"• {name}: <b>{grade}</b>")

    await send_target.answer("\n".join(lines), parse_mode="HTML")


@dp.message(Command("baholarim"))
async def baholarim_command(message: Message):
    await _send_baholarim(message, message.from_user.id)



    waiting_content = State()
    waiting_confirm = State()


def _broadcast_confirm_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="✅ Yuborish", callback_data="bcast_send"),
                InlineKeyboardButton(text="❌ Bekor qilish", callback_data="bcast_cancel"),
            ]
        ]
    )


@dp.message(Command("xabar"))
async def xabar_command(message: Message, state: FSMContext):
    if not is_admin(message.from_user.id):
        await message.answer("⛔️ Bu buyruq faqat admin uchun.")
        return
    await state.set_state(Broadcast.waiting_content)
    await message.answer(
        "📢 Barcha foydalanuvchilarga yuboriladigan xabarni yuboring "
        "(matn, rasm yoki hujjat bo'lishi mumkin).\n"
        "Bekor qilish uchun /bekor buyrug'ini yuboring."
    )


@dp.message(Broadcast.waiting_content)
async def broadcast_content_handler(message: Message, state: FSMContext):
    await state.update_data(_bcast_chat_id=message.chat.id, _bcast_message_id=message.message_id)
    await state.set_state(Broadcast.waiting_confirm)
    total = len(db_get_all_user_ids())
    await message.answer(
        f"⬆️ Xabar shu ko'rinishda <b>{total}</b> ta foydalanuvchiga yuboriladi. Tasdiqlaysizmi?",
        reply_markup=_broadcast_confirm_kb(),
        parse_mode="HTML",
    )


@dp.callback_query(Broadcast.waiting_confirm, F.data == "bcast_cancel")
async def broadcast_cancel_callback(call: CallbackQuery, state: FSMContext):
    await state.clear()
    await call.message.answer("❌ Xabar yuborish bekor qilindi.")
    await call.answer()


@dp.callback_query(Broadcast.waiting_confirm, F.data == "bcast_send")
async def broadcast_send_callback(call: CallbackQuery, state: FSMContext):
    if not is_admin(call.from_user.id):
        await call.answer("⛔️ Ruxsat yo'q.", show_alert=True)
        return
    data = await state.get_data()
    src_chat_id = data.get("_bcast_chat_id")
    src_message_id = data.get("_bcast_message_id")
    await state.clear()
    await call.answer()

    user_ids = db_get_all_user_ids()
    await call.message.answer(f"📤 Yuborish boshlandi... (jami {len(user_ids)} foydalanuvchi)")

    sent, failed = 0, 0
    for uid in user_ids:
        try:
            await bot.copy_message(chat_id=uid, from_chat_id=src_chat_id, message_id=src_message_id)
            sent += 1
        except Exception:
            failed += 1
        await asyncio.sleep(0.05)  # Telegram flood-limitiga tushib qolmaslik uchun

    await call.message.answer(
        f"✅ Xabar yuborildi.\n📨 Yetib bordi: <b>{sent}</b>\n🚫 Yetib bormadi: <b>{failed}</b>",
        parse_mode="HTML",
    )


class PaymentFlow(StatesGroup):
    waiting_receipt = State()


SERVICE_NAMES = {"3x4": "📸 3x4 rasm", "obyektivka": "🗂 Ma'lumotnoma"}
SERVICE_PRICES = {"3x4": PRICE_3X4, "obyektivka": PRICE_OBYEKTIVKA}


def _payment_prompt_kb(service: str) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="📤 To'lov chekini yuborish", callback_data=f"pay_start:{service}")],
        ]
    )


async def _send_payment_prompt(message: Message, service: str) -> None:
    """Xizmatdan foydalanishdan oldin to'lov haqida ma'lumot va chek yuborish
    tugmasini ko'rsatadi (faqat admin bo'lmagan foydalanuvchilar uchun).
    Natija (rasm yoki hujjat) faqat admin to'lovni tasdiqlagandan keyin yuboriladi."""
    price = SERVICE_PRICES[service]
    name = SERVICE_NAMES[service]
    await message.answer(
        "💳 <b>To'lov talab qilinadi</b>\n\n"
        f"Xizmat: {name}\n"
        f"Narxi: {price:,} so'm\n\n"
        "To'lovni quyidagi kartaga o'tkazing:\n"
        f"💳 <code>{PAYMENT_CARD_NUMBER}</code> ({PAYMENT_CARD_OWNER})\n\n"
        "To'lovni amalga oshirgandan so'ng, pastdagi tugma orqali chekni "
        "(skrinshot yoki rasm ko'rinishida) yuboring. Admin tekshirib "
        "tasdiqlagach, natijangiz (rasm/hujjat) sizga avtomatik yuboriladi.".replace(",", " "),
        reply_markup=_payment_prompt_kb(service),
        parse_mode="HTML",
    )


@dp.callback_query(F.data.startswith("pay_start:"))
async def pay_start_callback(call: CallbackQuery, state: FSMContext):
    service = call.data.split(":", 1)[1]
    if service not in SERVICE_PRICES:
        await call.answer()
        return

    user_id = call.from_user.id
    if service == "3x4" and user_id not in pending_3x4_payload:
        await call.answer(
            "⚠️ So'rovingiz muddati tugagan. Iltimos, 3x4 rasmni qaytadan yuboring.",
            show_alert=True,
        )
        return
    if service == "obyektivka" and user_id not in pending_obyektivka_payload:
        await call.answer(
            "⚠️ So'rovingiz muddati tugagan. Iltimos, /malumotnoma orqali qaytadan boshlang.",
            show_alert=True,
        )
        return

    await state.update_data(_payment_service=service)
    await state.set_state(PaymentFlow.waiting_receipt)
    await call.message.answer("📤 To'lov chekining skrinshotini (rasm ko'rinishida) yuboring:")
    await call.answer()


async def _handle_receipt(message: Message, state: FSMContext, file_id: str) -> None:
    data = await state.get_data()
    service = data.get("_payment_service", "")
    price = SERVICE_PRICES.get(service, 0)
    name = SERVICE_NAMES.get(service, service)
    user = message.from_user
    full_name = " ".join(filter(None, [user.first_name, user.last_name])) or "—"

    payment_id = db_create_payment_request(user.id, user.username, full_name, service, price, file_id)
    await state.clear()

    await message.answer(
        "✅ Chekingiz qabul qilindi, admin tomonidan tekshirilmoqda.\n"
        "Tasdiqlangach yoki rad etilgach, sizga xabar beriladi."
    )

    username_text = f"@{user.username}" if user.username else "username yo'q"
    caption = (
        "🧾 <b>Yangi to'lov cheki</b>\n\n"
        f"👤 Foydalanuvchi: {full_name} ({username_text})\n"
        f"🆔 ID: <code>{user.id}</code>\n"
        f"🛍 Xizmat: {name}\n"
        f"💵 Summasi: {price:,} so'm\n"
        f"🔖 So'rov raqami: #{payment_id}".replace(",", " ")
    )
    kb = InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="✅ Tasdiqlash", callback_data=f"pay_confirm:{payment_id}"),
                InlineKeyboardButton(text="❌ Rad etish", callback_data=f"pay_reject:{payment_id}"),
            ]
        ]
    )

    for admin_id in ADMIN_IDS | db_get_dynamic_admin_ids():
        try:
            await bot.send_photo(admin_id, file_id, caption=caption, reply_markup=kb, parse_mode="HTML")
        except Exception:
            logging.exception("Adminga (%s) chek yuborishda xatolik", admin_id)


@dp.message(PaymentFlow.waiting_receipt, F.photo)
async def receipt_photo_handler(message: Message, state: FSMContext):
    await _handle_receipt(message, state, message.photo[-1].file_id)


@dp.message(PaymentFlow.waiting_receipt, F.document)
async def receipt_document_handler(message: Message, state: FSMContext):
    if message.document.mime_type and message.document.mime_type.startswith("image/"):
        await _handle_receipt(message, state, message.document.file_id)
    else:
        await message.answer("Iltimos, chekni rasm (screenshot) ko'rinishida yuboring.")


@dp.message(PaymentFlow.waiting_receipt)
async def receipt_wrong_type(message: Message):
    await message.answer("Iltimos, to'lov chekini rasm ko'rinishida yuboring.")


@dp.callback_query(F.data.startswith("pay_confirm:"))
async def pay_confirm_callback(call: CallbackQuery):
    if not is_admin(call.from_user.id):
        await call.answer("⛔️ Ruxsat yo'q.", show_alert=True)
        return
    payment_id = int(call.data.split(":", 1)[1])
    payment = db_get_payment(payment_id)
    if payment is None:
        await call.answer("So'rov topilmadi.", show_alert=True)
        return
    if payment["status"] != "pending":
        await call.answer("Bu so'rov allaqachon ko'rib chiqilgan.", show_alert=True)
        return

    db_set_payment_status(payment_id, "confirmed", call.from_user.id)
    try:
        await call.message.edit_caption(
            caption=(call.message.caption or "") + f"\n\n✅ <b>Tasdiqlandi</b> (admin: {call.from_user.id})",
            parse_mode="HTML",
        )
    except Exception:
        pass

    target_user_id = payment["user_id"]
    service = payment["service"]
    try:
        if service == "3x4":
            raw_bytes = pending_3x4_payload.pop(target_user_id, None)
            if raw_bytes is None:
                await bot.send_message(
                    target_user_id,
                    "✅ To'lovingiz tasdiqlandi, lekin rasmingiz topilmadi. "
                    "Iltimos, 3x4 rasmni qaytadan yuboring.",
                )
            else:
                await bot.send_message(target_user_id, "✅ To'lovingiz tasdiqlandi! Natijangiz tayyorlanmoqda...")
                await _generate_and_send_3x4(target_user_id, target_user_id, raw_bytes)
        elif service == "obyektivka":
            payload = pending_obyektivka_payload.pop(target_user_id, None)
            if payload is None:
                await bot.send_message(
                    target_user_id,
                    "✅ To'lovingiz tasdiqlandi, lekin ma'lumotlaringiz topilmadi. "
                    "Iltimos, /malumotnoma orqali qaytadan boshlang.",
                )
            else:
                await bot.send_message(target_user_id, "✅ To'lovingiz tasdiqlandi! Hujjat tayyorlanmoqda...")
                await _generate_and_send_obyektivka(target_user_id, target_user_id, payload["data"], payload["photo_img"])
    except Exception:
        logging.exception("To'lov tasdiqlangandan keyin natija yuborishda xatolik")
    await call.answer("Tasdiqlandi ✅")


@dp.callback_query(F.data.startswith("pay_reject:"))
async def pay_reject_callback(call: CallbackQuery):
    if not is_admin(call.from_user.id):
        await call.answer("⛔️ Ruxsat yo'q.", show_alert=True)
        return
    payment_id = int(call.data.split(":", 1)[1])
    payment = db_get_payment(payment_id)
    if payment is None:
        await call.answer("So'rov topilmadi.", show_alert=True)
        return
    if payment["status"] != "pending":
        await call.answer("Bu so'rov allaqachon ko'rib chiqilgan.", show_alert=True)
        return

    db_set_payment_status(payment_id, "rejected", call.from_user.id)
    try:
        await call.message.edit_caption(
            caption=(call.message.caption or "") + f"\n\n❌ <b>Rad etildi</b> (admin: {call.from_user.id})",
            parse_mode="HTML",
        )
    except Exception:
        pass

    target_user_id = payment["user_id"]
    # Tasdiqlanmagani uchun kutilayotgan natija (rasm/ma'lumotlar) bekor qilinadi
    pending_3x4_payload.pop(target_user_id, None)
    pending_obyektivka_payload.pop(target_user_id, None)
    try:
        await bot.send_message(
            target_user_id,
            "❌ To'lov chekingiz rad etildi.\n"
            "Iltimos, to'g'ri va aniq chekni yuboring yoki admin bilan bog'lanib, "
            "xizmatdan (3x4 rasm yoki /malumotnoma) qaytadan foydalaning.",
        )
    except Exception:
        logging.info("Foydalanuvchiga rad etish haqida xabar yuborib bo'lmadi.")
    await call.answer("Rad etildi ❌")


async def _generate_and_send_3x4(user_id: int, chat_id: int, image_bytes: bytes) -> None:
    """3x4 rasmni tayyorlab, foydalanuvchiga yuboradi. To'lov talab qilinmaydigan
    (admin) yoki to'lov allaqachon tasdiqlangan hollarda chaqiriladi."""
    try:
        cropped = make_3x4(image_bytes)
        white_bg = remove_bg_to_white(cropped)
        pending_photos[user_id] = white_bg

        await bot.send_document(
            chat_id,
            BufferedInputFile(_to_jpeg_bytes(white_bg), filename="3x4.jpg"),
            caption="Tayyor ✅ 3x4 (354x472 px, 300 DPI), fon oqartirildi.",
            reply_markup=_print_sheet_kb(),
        )
    except FaceNotFoundError:
        await bot.send_message(chat_id, "😕 Yuz aniq ko'rinmadi, boshqa rasm yuboring.")
    except Exception as e:
        logging.exception("Xatolik")
        db_log_error(user_id, str(e))
        await bot.send_message(chat_id, f"Kechirasiz, xatolik yuz berdi: {e}")


async def _process_and_reply(message: Message, image_bytes: bytes):
    user_id = message.from_user.id

    if _is_low_quality(image_bytes):
        await message.answer(
            "⚠️ Rasm juda kichik yoki loyqa ko'rinmoqda.\n"
            "Iltimos, sifatliroq (aniq va yetarlicha o'lchamdagi) rasm yuboring."
        )
        return

    if is_admin(user_id):
        await _generate_and_send_3x4(user_id, message.chat.id, image_bytes)
        return

    # Admin bo'lmagan foydalanuvchilar uchun: avval to'lov so'raladi, natija
    # faqat admin chekni tasdiqlagandan keyin yuboriladi.
    pending_3x4_payload[user_id] = image_bytes
    await _send_payment_prompt(message, "3x4")


@dp.callback_query(F.data == "print_sheet")
async def print_sheet_callback(call: CallbackQuery):
    photo = pending_photos.get(call.from_user.id)
    if photo is None:
        await call.answer("Avval rasm yuboring.", show_alert=True)
        return

    sheet = make_print_sheet(photo)
    await call.message.answer_document(
        BufferedInputFile(_to_jpeg_bytes(sheet), filename="3x4_varaq.jpg"),
        caption="Chop etish varag'i tayyor ✅ (8 dona 3x4)",
    )
    await call.answer()


@dp.message(F.photo, StateFilter(None))
async def photo_handler(message: Message):
    photo = message.photo[-1]
    file = await bot.get_file(photo.file_id)
    file_bytes = await bot.download_file(file.file_path)
    await _process_and_reply(message, file_bytes.read())


@dp.message(F.document, StateFilter(None))
async def document_handler(message: Message):
    if message.document.mime_type and message.document.mime_type.startswith("image/"):
        file = await bot.get_file(message.document.file_id)
        file_bytes = await bot.download_file(file.file_path)
        await _process_and_reply(message, file_bytes.read())
    else:
        await message.answer("Iltimos, rasm fayl yuboring.")



# ============================================================
# OBYEKTIVKA (MA'LUMOTNOMA) BO'LIMI
# ============================================================

# MUHIM: bu handler pastdagi barcha FSM (Obyektivka.*) handlerlaridan OLDIN
# ro'yxatga olinishi shart — aiogram handlerlarni ro'yxatga olingan tartibda
# tekshiradi, shuning uchun /bekor har qanday savol bosqichida ham ishlashi
# uchun state-ga bog'liq handlerlardan yuqorida turishi kerak.
@dp.message(Command("bekor"))
async def cancel_handler(message: Message, state: FSMContext):
    current_state = await state.get_state()
    if current_state is None:
        await message.answer("Hozircha faol jarayon yo'q.")
        return
    pending_confirm_photo.pop(message.from_user.id, None)
    await state.clear()
    await message.answer(
        "❌ Jarayon bekor qilindi.\n"
        "Qaytadan boshlash uchun /malumotnoma buyrug'ini yuboring."
    )


@dp.message(F.text == "❌ Bekor qilish")
async def menu_cancel_button(message: Message, state: FSMContext):
    await cancel_handler(message, state)


class Obyektivka(StatesGroup):
    full_name = State()
    birth_date = State()
    birth_place = State()
    nationality = State()
    party = State()
    education_level = State()
    institution = State()
    specialty = State()
    academic_degree = State()
    academic_title = State()
    languages = State()
    awards = State()
    deputy_status = State()
    work_period = State()
    work_description = State()
    work_more = State()
    relative_relation = State()
    relative_fish = State()
    relative_birth = State()
    relative_work = State()
    relative_address = State()
    relative_more = State()
    phone = State()
    waiting_photo = State()
    confirm = State()


def _sample_photo() -> Image.Image:
    """Namuna uchun oddiy siluet rasm (haqiqiy foto emas)."""
    w, h = TARGET_WIDTH, TARGET_HEIGHT
    img = Image.new("RGB", (w, h), (235, 235, 235))
    draw = ImageDraw.Draw(img)
    cx = w // 2
    head_r = int(w * 0.20)
    head_cy = int(h * 0.32)
    draw.ellipse(
        [cx - head_r, head_cy - head_r, cx + head_r, head_cy + head_r],
        fill=(170, 170, 170),
    )
    draw.polygon(
        [
            (cx - int(w * 0.42), h),
            (cx + int(w * 0.42), h),
            (cx + int(w * 0.26), int(h * 0.60)),
            (cx - int(w * 0.26), int(h * 0.60)),
        ],
        fill=(170, 170, 170),
    )
    return img


def _sample_obyektivka_data() -> dict:
    """Namuna (misol) uchun to'liq o'ylab topilgan, hech kimga tegishli
    bo'lmagan namunaviy ma'lumotlar."""
    return {
        "full_name": "Falonchiyev Falon Falonovich",
        "birth_date": "01.01.2000-yil",
        "birth_place": "Toshkent shahri",
        "nationality": "o'zbek",
        "party": "Partiyasiz",
        "education_level": "Oliy",
        "institution": "Namuna Davlat Universiteti",
        "specialty": "Namuna yo'nalishi",
        "academic_degree": "yo'q",
        "academic_title": "yo'q",
        "languages": "ingliz tili (o'rta darajada)",
        "awards": "yo'q",
        "deputy_status": "yo'q",
        "work_entries": [
            {"period": "2018-2022 yy", "description": "Namuna Davlat Universiteti talabasi"},
            {"period": "2022 y - hozirgacha", "description": "\"Namuna\" MChJ, mutaxassis"},
        ],
        "relatives": [
            {
                "relation": "Otasi",
                "fish": "Falonchiyev Falon Falonovich",
                "birth": "1975-yil Toshkent shahri",
                "work": "\"Namuna\" korxonasi, muhandis",
                "address": "Toshkent shahri, Namuna ko'chasi",
            },
            {
                "relation": "Onasi",
                "fish": "Falonchiyeva Falona Falonovna",
                "birth": "1978-yil Toshkent shahri",
                "work": "Uy bekasi",
                "address": "Toshkent shahri, Namuna ko'chasi",
            },
        ],
        "phone": "901234567",
    }


async def _send_sample_document(message: Message):
    sample_docx = _build_obyektivka_docx(_sample_obyektivka_data(), _sample_photo())
    await message.answer_document(
        BufferedInputFile(sample_docx, filename="namuna_malumotnoma.docx"),
        caption=(
            "📄 Bu — namuna (barcha F.I.SH. va ma'lumotlar o'ylab topilgan, "
            "haqiqiy emas). Sizning ma'lumotnomangiz aynan shu ko'rinishda, "
            "lekin siz kiritgan haqiqiy ma'lumotlar va rasmingiz bilan tayyor bo'ladi."
        ),
    )


@dp.message(Command("namuna"))
async def namuna_handler(message: Message):
    await _send_sample_document(message)


@dp.message(F.text == "📋 Namuna")
async def menu_namuna_button(message: Message):
    await namuna_handler(message)


@dp.message(F.text == "📊 Holatim")
async def menu_holatim_button(message: Message):
    await holatim_command(message)


@dp.message(F.text == "🎓 HEMIS")
async def menu_hemis_button(message: Message, state: FSMContext):
    await hemis_command(message, state)


@dp.message(F.text == "❓ Yordam")
async def menu_yordam_button(message: Message):
    await yordam_command(message)


def _more_relatives_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="➕ Yana qo'shish", callback_data="rel_more")],
            [InlineKeyboardButton(text="✅ Tugatish", callback_data="rel_done")],
        ]
    )


def _more_work_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="➕ Yana qo'shish", callback_data="work_more")],
            [InlineKeyboardButton(text="✅ Tugatish", callback_data="work_done")],
        ]
    )


# Vaqtinchalik xotira: user_id -> obyektivka uchun tayyorlangan (fon oqartirilgan)
# rasm, tasdiqlash bosqichida kutib turadi.
pending_confirm_photo: dict[int, Image.Image] = {}


async def _start_malumotnoma_flow(send_target, user_id: int, state: FSMContext) -> None:
    """send_target — javob yuborish uchun ishlatiladigan obyekt (Message yoki
    CallbackQuery.message), user_id — jarayon boshlanayotgan haqiqiy foydalanuvchi."""
    pending_confirm_photo.pop(user_id, None)
    await state.clear()
    await send_target.answer("⏳ Avval namuna ko'rinishini yuboraman...")
    await _send_sample_document(send_target)

    await state.set_data({"relatives": [], "work_entries": []})
    await state.set_state(Obyektivka.full_name)
    await send_target.answer(
        "Endi sizning haqiqiy ma'lumotnomangizni tuzishni boshlaymiz.\n\n"
        "Istalgan vaqtda /bekor buyrug'i bilan jarayonni to'xtatishingiz mumkin.\n\n"
        "To'liq FIO (Familiya Ism Sharifingiz) kiriting:"
    )


@dp.message(Command("malumotnoma"))
async def malumotnoma_start(message: Message, state: FSMContext):
    await _start_malumotnoma_flow(message, message.from_user.id, state)


@dp.message(F.text == "🗂 Ma'lumotnoma")
async def menu_malumotnoma_button(message: Message, state: FSMContext):
    await malumotnoma_start(message, state)




@dp.message(Obyektivka.full_name)
async def q_full_name(message: Message, state: FSMContext):
    await state.update_data(full_name=message.text)
    await state.set_state(Obyektivka.birth_date)
    await message.answer("Tug'ilgan yilingizni kiriting (masalan: 15.03.2000-yil):")


@dp.message(Obyektivka.birth_date)
async def q_birth_date(message: Message, state: FSMContext):
    if not is_valid_date(message.text or ""):
        await message.answer(
            "❗️ Sana noto'g'ri formatda kiritildi.\n"
            "Iltimos, 'kun.oy.yil' ko'rinishida kiriting (masalan: 15.03.2000):"
        )
        return
    await state.update_data(birth_date=message.text)
    await state.set_state(Obyektivka.birth_place)
    await message.answer("Tug'ilgan joyingizni kiriting (viloyat, tuman):")


@dp.message(Obyektivka.birth_place)
async def q_birth_place(message: Message, state: FSMContext):
    await state.update_data(birth_place=message.text)
    await state.set_state(Obyektivka.nationality)
    await message.answer("Millatingizni kiriting:")


@dp.message(Obyektivka.nationality)
async def q_nationality(message: Message, state: FSMContext):
    await state.update_data(nationality=message.text)
    await state.set_state(Obyektivka.party)
    await message.answer("Partiyaviyligini kiriting (masalan: Partiyasiz):")


@dp.message(Obyektivka.party)
async def q_party(message: Message, state: FSMContext):
    await state.update_data(party=message.text)
    await state.set_state(Obyektivka.education_level)
    await message.answer("Ma'lumotingizni kiriting (Oliy, o'rta-maxsus, o'rta va h.k.):")


@dp.message(Obyektivka.education_level)
async def q_education_level(message: Message, state: FSMContext):
    await state.update_data(education_level=message.text)
    await state.set_state(Obyektivka.institution)
    await message.answer("Ta'lim muassasasini (universitet/maktab nomi) kiriting:")


@dp.message(Obyektivka.institution)
async def q_institution(message: Message, state: FSMContext):
    await state.update_data(institution=message.text)
    await state.set_state(Obyektivka.specialty)
    await message.answer("Mutaxassisligingizni kiriting:")


@dp.message(Obyektivka.specialty)
async def q_specialty(message: Message, state: FSMContext):
    await state.update_data(specialty=message.text)
    await state.set_state(Obyektivka.academic_degree)
    await message.answer("Ilmiy darajangiz bormi? (Yo'q bo'lsa 'yo'q' deb yozing):")


@dp.message(Obyektivka.academic_degree)
async def q_academic_degree(message: Message, state: FSMContext):
    await state.update_data(academic_degree=message.text)
    await state.set_state(Obyektivka.academic_title)
    await message.answer("Ilmiy unvoningiz bormi? (Yo'q bo'lsa 'yo'q' deb yozing):")


@dp.message(Obyektivka.academic_title)
async def q_academic_title(message: Message, state: FSMContext):
    await state.update_data(academic_title=message.text)
    await state.set_state(Obyektivka.languages)
    await message.answer("Qaysi chet tillarini bilasiz?")


@dp.message(Obyektivka.languages)
async def q_languages(message: Message, state: FSMContext):
    await state.update_data(languages=message.text)
    await state.set_state(Obyektivka.awards)
    await message.answer("Davlat mukofotlari bilan taqdirlanganmisiz? (Qanaqa / Yo'q):")


@dp.message(Obyektivka.awards)
async def q_awards(message: Message, state: FSMContext):
    await state.update_data(awards=message.text)
    await state.set_state(Obyektivka.deputy_status)
    await message.answer(
        "Xalq deputatlari Kengashi deputatimisiz yoki boshqa saylanadigan "
        "organ a'zosi? (To'liq yozing):"
    )


@dp.message(Obyektivka.deputy_status)
async def q_deputy_status(message: Message, state: FSMContext):
    await state.update_data(deputy_status=message.text)
    await state.set_state(Obyektivka.work_period)
    await message.answer(
        "Endi **MEHNAT FAOLIYATI**ni kiritamiz.\n\n"
        "Davrini kiriting (masalan: 2015-2020 yy):"
    )


@dp.message(Obyektivka.work_period)
async def q_work_period(message: Message, state: FSMContext):
    await state.update_data(_current_work_period=message.text)
    await state.set_state(Obyektivka.work_description)
    await message.answer(
        "Shu davrdagi ish/o'qish joyi va lavozimingizni (yoki faoliyatingizni) kiriting "
        "(masalan: 20-maktab o'quvchisi):"
    )


@dp.message(Obyektivka.work_description)
async def q_work_description(message: Message, state: FSMContext):
    data = await state.get_data()
    work_entries = data.get("work_entries", [])
    work_entries.append({"period": data.get("_current_work_period", ""), "description": message.text})
    await state.update_data(work_entries=work_entries)
    await state.set_state(Obyektivka.work_more)
    await message.answer(
        "Yana mehnat faoliyati davri qo'shasizmi?",
        reply_markup=_more_work_kb(),
    )


@dp.callback_query(Obyektivka.work_more, F.data == "work_more")
async def work_more(call: CallbackQuery, state: FSMContext):
    await state.set_state(Obyektivka.work_period)
    await call.message.answer("Davrini kiriting (masalan: 2020 y h.v):")
    await call.answer()


@dp.callback_query(Obyektivka.work_more, F.data == "work_done")
async def work_done(call: CallbackQuery, state: FSMContext):
    await state.set_state(Obyektivka.relative_relation)
    await call.message.answer(
        "Endi **YAQIN QARINDOSHLARI HAQIDA MA'LUMOT**ni kiritamiz.\n\n"
        "Qarindoshlik darajasini kiriting (masalan: Otasi, Onasi...):"
    )
    await call.answer()


@dp.message(Obyektivka.relative_relation)
async def q_relative_relation(message: Message, state: FSMContext):
    await state.update_data(_current_relation=message.text)
    await state.set_state(Obyektivka.relative_fish)
    await message.answer("Shu qarindoshingizning Familiyasi, ismi va otasining ismini kiriting:")


@dp.message(Obyektivka.relative_fish)
async def q_relative_fish(message: Message, state: FSMContext):
    await state.update_data(_current_relative_fish=message.text)
    await state.set_state(Obyektivka.relative_birth)
    await message.answer("Tug'ilgan yili va joyini kiriting (masalan: 1975-yil Toshkent shahri):")


@dp.message(Obyektivka.relative_birth)
async def q_relative_birth(message: Message, state: FSMContext):
    await state.update_data(_current_relative_birth=message.text)
    await state.set_state(Obyektivka.relative_work)
    await message.answer("Ish joyi va lavozimini kiriting:")


@dp.message(Obyektivka.relative_work)
async def q_relative_work(message: Message, state: FSMContext):
    await state.update_data(_current_relative_work=message.text)
    await state.set_state(Obyektivka.relative_address)
    await message.answer("Turar joyini (manzilini) kiriting:")


@dp.message(Obyektivka.relative_address)
async def q_relative_address(message: Message, state: FSMContext):
    data = await state.get_data()
    relatives = data.get("relatives", [])
    relatives.append({
        "relation": data.get("_current_relation", ""),
        "fish": data.get("_current_relative_fish", ""),
        "birth": data.get("_current_relative_birth", ""),
        "work": data.get("_current_relative_work", ""),
        "address": message.text,
    })
    await state.update_data(relatives=relatives)
    await state.set_state(Obyektivka.relative_more)
    await message.answer(
        "Yana qarindosh qo'shasizmi?",
        reply_markup=_more_relatives_kb(),
    )


@dp.callback_query(Obyektivka.relative_more, F.data == "rel_more")
async def relative_more(call: CallbackQuery, state: FSMContext):
    await state.set_state(Obyektivka.relative_relation)
    await call.message.answer("Qarindoshlik darajasini kiriting (masalan: Aka, Uka, Opa...):")
    await call.answer()


@dp.callback_query(Obyektivka.relative_more, F.data == "rel_done")
async def relative_done(call: CallbackQuery, state: FSMContext):
    await state.set_state(Obyektivka.phone)
    await call.message.answer("Aloqa uchun telefon raqamingizni kiriting (masalan: 901234567):")
    await call.answer()


@dp.message(Obyektivka.phone)
async def q_phone(message: Message, state: FSMContext):
    if not is_valid_phone(message.text or ""):
        await message.answer(
            "❗️ Telefon raqami noto'g'ri kiritildi.\n"
            "Raqam faqat sonlardan iborat va 9-13 xonali bo'lishi kerak "
            "(masalan: 901234567). Qaytadan kiriting:"
        )
        return
    await state.update_data(phone=message.text)
    await state.set_state(Obyektivka.waiting_photo)
    await message.answer("Endi **3x4 formatdagi shaxsiy rasmingizni** yuboring:")


def _format_summary(data: dict) -> str:
    """Foydalanuvchi kiritgan barcha ma'lumotlarni tasdiqlash uchun matn ko'rinishida yig'adi."""
    lines = ["📋 <b>Kiritilgan ma'lumotlaringiz:</b>\n"]
    lines.append(f"👤 FIO: {data.get('full_name', '')}")
    lines.append(f"🎂 Tug'ilgan sana: {data.get('birth_date', '')}")
    lines.append(f"📍 Tug'ilgan joyi: {data.get('birth_place', '')}")
    lines.append(f"🌍 Millati: {data.get('nationality', '')}")
    lines.append(f"🏛 Partiyaviyligi: {data.get('party', '')}")
    lines.append(f"🎓 Ma'lumoti: {data.get('education_level', '')}")
    lines.append(f"🏫 Tamomlagan: {data.get('institution', '')}")
    lines.append(f"🧪 Mutaxassisligi: {data.get('specialty', '')}")
    lines.append(f"🎖 Ilmiy darajasi: {data.get('academic_degree', '')}")
    lines.append(f"🎖 Ilmiy unvoni: {data.get('academic_title', '')}")
    lines.append(f"🗣 Tillar: {data.get('languages', '')}")
    lines.append(f"🏅 Mukofotlar: {data.get('awards', '')}")
    lines.append(f"🏛 Deputatlik: {data.get('deputy_status', '')}")

    work_entries = data.get("work_entries", [])
    if work_entries:
        lines.append("\n💼 <b>Mehnat faoliyati:</b>")
        for w in work_entries:
            lines.append(f"  • {w.get('period', '')} — {w.get('description', '')}")

    relatives = data.get("relatives", [])
    if relatives:
        lines.append("\n👪 <b>Qarindoshlar:</b>")
        for r in relatives:
            lines.append(f"  • {r.get('relation', '')}: {r.get('fish', '')}")

    lines.append(f"\n📞 Telefon: {data.get('phone', '')}")
    lines.append("\n❓ <b>Hammasi to'g'rimi?</b>")
    return "\n".join(lines)


def _confirm_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="✅ To'g'ri, hujjat yaratilsin", callback_data="confirm_yes")],
            [InlineKeyboardButton(text="🔄 Qaytadan boshlash", callback_data="confirm_restart")],
        ]
    )


def _find_soffice_executable() -> str:
    """soffice (LibreOffice) bajariluvchi faylini turli manbalardan qidiradi."""
    import shutil

    # 1) Muhit o'zgaruvchisi orqali aniq ko'rsatilgan bo'lsa
    env_path = os.getenv("SOFFICE_PATH")
    if env_path and os.path.isfile(env_path):
        return env_path

    # 2) PATH ichidan qidirish (Linux/macOS'da odatda shu yetarli)
    found = shutil.which("soffice") or shutil.which("soffice.exe") or shutil.which("soffice.bin")
    if found:
        return found

    # 3) Windows'dagi standart o'rnatish joylari
    if os.name == "nt":
        candidates = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
        for path in candidates:
            if os.path.isfile(path):
                return path

    raise RuntimeError(
        "LibreOffice (soffice) topilmadi. DOCX'ni PDF'ga aylantirish uchun "
        "kompyuteringizga LibreOffice o'rnatilgan bo'lishi kerak "
        "(https://www.libreoffice.org/download/download/).\n\n"
        "O'rnatgandan so'ng ham xatolik davom etsa, muhit o'zgaruvchisi orqali "
        "to'liq yo'lni ko'rsating, masalan Windows'da:\n"
        r'SOFFICE_PATH=C:\Program Files\LibreOffice\program\soffice.exe'
    )


async def convert_docx_to_pdf(docx_bytes: bytes) -> bytes:
    """LibreOffice (soffice) yordamida DOCX faylni PDF'ga aylantiradi."""
    soffice_exe = _find_soffice_executable()

    with tempfile.TemporaryDirectory() as tmp_dir:
        docx_path = os.path.join(tmp_dir, "malumotnoma.docx")
        with open(docx_path, "wb") as f:
            f.write(docx_bytes)

        try:
            proc = await asyncio.create_subprocess_exec(
                soffice_exe, "--headless", "--norestore",
                "--convert-to", "pdf", "--outdir", tmp_dir, docx_path,
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.PIPE,
            )
        except (FileNotFoundError, OSError) as e:
            raise RuntimeError(
                f"LibreOffice'ni ishga tushirib bo'lmadi ({soffice_exe}): {e}\n"
                "LibreOffice to'g'ri o'rnatilganini tekshiring."
            )

        try:
            stdout, stderr = await asyncio.wait_for(proc.communicate(), timeout=90)
        except asyncio.TimeoutError:
            proc.kill()
            raise RuntimeError("PDF'ga aylantirish vaqti tugadi (LibreOffice javob bermadi).")

        pdf_path = os.path.join(tmp_dir, "malumotnoma.pdf")
        if not os.path.isfile(pdf_path):
            err_text = stderr.decode(errors="ignore") if stderr else "noma'lum xato"
            raise RuntimeError(f"PDF konvertatsiyasida xatolik: {err_text}")

        with open(pdf_path, "rb") as f:
            return f.read()


@dp.message(Obyektivka.waiting_photo, F.photo)
async def q_photo(message: Message, state: FSMContext):
    user_id = message.from_user.id
    try:
        photo = message.photo[-1]
        file = await bot.get_file(photo.file_id)
        file_bytes = await bot.download_file(file.file_path)
        raw_bytes = file_bytes.read()

        if _is_low_quality(raw_bytes):
            await message.answer(
                "⚠️ Rasm juda kichik yoki loyqa ko'rinmoqda.\n"
                "Iltimos, sifatliroq (aniq va yetarlicha o'lchamdagi) rasm yuboring."
            )
            return

        cropped = make_3x4(raw_bytes)
        photo_img = remove_bg_to_white(cropped)
    except FaceNotFoundError:
        await message.answer("😕 Yuz aniq ko'rinmadi, boshqa rasm yuboring.")
        return
    except Exception as e:
        logging.exception("Obyektivka rasm xatoligi")
        db_log_error(user_id, str(e))
        await message.answer(f"Kechirasiz, xatolik yuz berdi: {e}")
        return

    pending_confirm_photo[user_id] = photo_img
    data = await state.get_data()
    await state.set_state(Obyektivka.confirm)
    await message.answer(_format_summary(data), reply_markup=_confirm_kb(), parse_mode="HTML")


@dp.message(Obyektivka.waiting_photo)
async def q_photo_wrong_type(message: Message):
    await message.answer("Iltimos, rasm (foto) ko'rinishida yuboring.")


async def _generate_and_send_obyektivka(user_id: int, chat_id: int, data: dict, photo_img: Image.Image) -> None:
    """Ma'lumotnoma hujjatini tuzib, foydalanuvchiga yuboradi. To'lov talab
    qilinmaydigan (admin) yoki to'lov allaqachon tasdiqlangan hollarda chaqiriladi."""
    try:
        docx_bytes = _build_obyektivka_docx(data, photo_img)
        pdf_bytes = await convert_docx_to_pdf(docx_bytes)

        # Ikkala faylni (.docx va .pdf) bitta media-guruh sifatida, birga yuboramiz
        media_group = [
            InputMediaDocument(
                media=BufferedInputFile(docx_bytes, filename="malumotnoma.docx"),
                caption="Ma'lumotnomangiz tayyor ✅ (.docx va .pdf)",
            ),
            InputMediaDocument(
                media=BufferedInputFile(pdf_bytes, filename="malumotnoma.pdf"),
            ),
        ]
        await bot.send_media_group(chat_id, media=media_group)
        db_log_request(user_id)
    except Exception as e:
        logging.exception("Hujjat yaratishda xatolik")
        db_log_error(user_id, str(e))
        await bot.send_message(chat_id, f"Kechirasiz, xatolik yuz berdi: {e}")


@dp.callback_query(Obyektivka.confirm, F.data == "confirm_yes")
async def confirm_yes(call: CallbackQuery, state: FSMContext):
    user_id = call.from_user.id
    photo_img = pending_confirm_photo.get(user_id)
    if photo_img is None:
        await call.answer("Rasm topilmadi, iltimos qaytadan yuboring.", show_alert=True)
        return

    data = await state.get_data()
    await call.answer()
    pending_confirm_photo.pop(user_id, None)
    await state.clear()

    if is_admin(user_id):
        await call.message.answer("⏳ Hujjat (.docx va .pdf) tayyorlanmoqda, biroz kuting...")
        await _generate_and_send_obyektivka(user_id, call.message.chat.id, data, photo_img)
        return

    # Admin bo'lmagan foydalanuvchilar uchun: avval to'lov so'raladi, hujjat
    # faqat admin chekni tasdiqlagandan keyin tuzilib yuboriladi.
    pending_obyektivka_payload[user_id] = {"data": data, "photo_img": photo_img}
    await _send_payment_prompt(call.message, "obyektivka")


@dp.callback_query(Obyektivka.confirm, F.data == "confirm_restart")
async def confirm_restart(call: CallbackQuery, state: FSMContext):
    user_id = call.from_user.id
    pending_confirm_photo.pop(user_id, None)
    await call.answer()

    await state.set_data({"relatives": [], "work_entries": []})
    await state.set_state(Obyektivka.full_name)
    await call.message.answer(
        "🔄 Qaytadan boshlaymiz.\n\nTo'liq FIO (Familiya Ism Sharifingiz) kiriting:"
    )


FONT_NAME = "Times New Roman"


def _set_run_font(run, size_pt=11, bold=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.bold = bold
    # Sharq (kirill/lotin) shriftlar to'g'ri ko'rsatilishi uchun
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), FONT_NAME)
    rFonts.set(qn("w:hAnsi"), FONT_NAME)
    rFonts.set(qn("w:eastAsia"), FONT_NAME)
    rFonts.set(qn("w:cs"), FONT_NAME)


def _add_paragraph(doc, text="", size=11, bold=False, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    _set_run_font(run, size_pt=size, bold=bold)
    return p, run


def _set_no_borders(table):
    """Jadval chiziqlarini butunlay yashiradi (obyektivka yuqori qismidagi
    juftlik-maydonlar ko'rinishidagi jadval kabi)."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def _add_floating_photo(paragraph, img_buf, width_cm=3.0, height_cm=4.0):
    """Rasmni sahifaning o'ng yuqori burchagiga, matn oqishi bilan
    (wrap square) joylaydi — shablondagi kabi."""
    run = paragraph.add_run()
    run.add_picture(img_buf, width=Cm(width_cm), height=Cm(height_cm))
    # add_picture inline rasm qo'shadi (w:drawing > wp:inline); buni
    # o'ng yuqori burchakka moslashtirilgan "floating" (wp:anchor) ga aylantiramiz.
    inline = run._element.find(qn("w:drawing")).find(qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    docPr = inline.find(qn("wp:docPr"))
    graphic = inline.find(qn("a:graphic"))

    cx = extent.get("cx")
    cy = extent.get("cy")

    anchor = OxmlElement("wp:anchor")
    for attr, val in {
        "distT": "0", "distB": "0", "distL": "114300", "distR": "114300",
        "simplePos": "0", "relativeHeight": "251658240", "behindDoc": "0",
        "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(attr, val)

    simplePos = OxmlElement("wp:simplePos")
    simplePos.set("x", "0")
    simplePos.set("y", "0")
    anchor.append(simplePos)

    positionH = OxmlElement("wp:positionH")
    positionH.set("relativeFrom", "margin")
    alignH = OxmlElement("wp:align")
    alignH.text = "right"
    positionH.append(alignH)
    anchor.append(positionH)

    positionV = OxmlElement("wp:positionV")
    positionV.set("relativeFrom", "margin")
    posOffsetV = OxmlElement("wp:posOffset")
    posOffsetV.text = "0"
    positionV.append(posOffsetV)
    anchor.append(positionV)

    new_extent = OxmlElement("wp:extent")
    new_extent.set("cx", cx)
    new_extent.set("cy", cy)
    anchor.append(new_extent)

    effectExtent = OxmlElement("wp:effectExtent")
    for a in ("l", "t", "r", "b"):
        effectExtent.set(a, "0")
    anchor.append(effectExtent)

    wrapSquare = OxmlElement("wp:wrapSquare")
    wrapSquare.set("wrapText", "bothSides")
    anchor.append(wrapSquare)

    anchor.append(docPr)
    cNvGraphicFramePr = inline.find(qn("wp:cNvGraphicFramePr"))
    if cNvGraphicFramePr is not None:
        anchor.append(cNvGraphicFramePr)
    anchor.append(graphic)

    drawing = run._element.find(qn("w:drawing"))
    drawing.remove(inline)
    drawing.append(anchor)


def _add_pair_row(table, row_idx, label1, value1, label2, value2):
    row = table.rows[row_idx]
    for cell, label, value in ((row.cells[0], label1, value1), (row.cells[1], label2, value2)):
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        if label:
            r = cell.paragraphs[0].add_run(f"{label}:")
            _set_run_font(r, size_pt=11, bold=True)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(8)
        r2 = p2.add_run(str(value) if value else "")
        _set_run_font(r2, size_pt=11, bold=False)


def _add_qa_block(doc, question, answer):
    _add_paragraph(doc, question, size=11, bold=False, space_after=0)
    _add_paragraph(doc, str(answer) if answer else "", size=11, bold=False, space_after=10)


def _build_obyektivka_docx(data: dict, photo_img: Image.Image) -> bytes:
    """Yig'ilgan ma'lumotlar va rasm asosida, shablonga mos Word hujjat yaratadi."""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(1.5)

    # Hujjat bo'yicha standart shrift
    normal_style = doc.styles["Normal"]
    normal_style.font.name = FONT_NAME
    normal_style.font.size = Pt(11)

    full_name = data.get("full_name", "")

    # ---------- SARLAVHA + RASM (o'ng yuqori burchakda) ----------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("MA'LUMOTNOMA")
    _set_run_font(title_run, size_pt=14, bold=True)

    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(full_name)
    _set_run_font(name_run, size_pt=14, bold=True)

    img_buf = io.BytesIO()
    frame_photo(photo_img).save(img_buf, format="JPEG", quality=95)
    img_buf.seek(0)
    _add_floating_photo(title_p, img_buf, width_cm=3.0, height_cm=4.0)

    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # ---------- ASOSIY MA'LUMOTLAR (chegarasiz 2 ustunli jadval) ----------
    info_table = doc.add_table(rows=5, cols=2)
    info_table.autofit = True
    _set_no_borders(info_table)

    _add_pair_row(info_table, 0, "Tug'ilgan yili", data.get("birth_date", ""),
                  "Tug'ilgan joyi", data.get("birth_place", ""))
    _add_pair_row(info_table, 1, "Millati", data.get("nationality", ""),
                  "Partiyaviyligi", data.get("party", ""))
    _add_pair_row(info_table, 2, "Ma'lumoti", data.get("education_level", ""),
                  "Tamomlagan", data.get("institution", ""))
    _add_pair_row(info_table, 3, "Ma'lumoti bo'yicha mutaxasisligi", data.get("specialty", ""),
                  "", "")
    _add_pair_row(info_table, 4, "Ilmiy darajasi", data.get("academic_degree", ""),
                  "Ilmiy unvoni", data.get("academic_title", ""))

    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    _add_qa_block(doc, "Qaysi chet va MDH larining tilini biladi (to'liq ko'rsatilsin):",
                  data.get("languages", ""))
    _add_qa_block(doc, "Davlat mukofotlari bilan taqdirlanganmi (qanaqa):",
                  data.get("awards", ""))
    _add_qa_block(
        doc,
        "Xalq deputatlari, respublika, viloyat, shahar va tuman Kengashi deputatimi "
        "yoki boshqa saylanadigan organlarning a'zosimi (to'liq korsatilishi lozim):",
        data.get("deputy_status", ""),
    )

    _add_paragraph(doc, "MEHNAT FAOLIYATI:", size=12, bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    for entry in data.get("work_entries", []):
        period = entry.get("period", "")
        desc = entry.get("description", "")
        _add_paragraph(doc, f"{period}\t{desc}", size=11, bold=False, space_after=4)

    # ---------- 2-SAHIFA: QARINDOSHLAR ----------
    doc.add_page_break()

    rel_title_p = doc.add_paragraph()
    rel_title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rel_title_p.paragraph_format.space_after = Pt(0)
    r1 = rel_title_p.add_run(full_name)
    _set_run_font(r1, size_pt=13, bold=True)
    r2 = rel_title_p.add_run("ning yaqin qarindoshlari haqida")
    _set_run_font(r2, size_pt=13, bold=True)

    _add_paragraph(doc, "MA'LUMOT", size=13, bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    headers = ["Qarindoshligi", "Familiyasi, ismi va otasining ismi",
               "Tug'ilgan yili va joyi", "Ish joyi va lovozimi", "Turar joyi"]
    relatives = data.get("relatives", [])

    rel_table = doc.add_table(rows=1 + max(len(relatives), 0), cols=5)
    rel_table.style = "Table Grid"
    rel_table.autofit = True

    header_cells = rel_table.rows[0].cells
    for i, h in enumerate(headers):
        header_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_cells[i].vertical_alignment = 1  # center
        run = header_cells[i].paragraphs[0].add_run(h)
        _set_run_font(run, size_pt=12, bold=True)

    for row_i, rel in enumerate(relatives, start=1):
        row_values = [
            rel.get("relation", ""), rel.get("fish", ""), rel.get("birth", ""),
            rel.get("work", ""), rel.get("address", ""),
        ]
        cells = rel_table.rows[row_i].cells
        for i, val in enumerate(row_values):
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cells[i].paragraphs[0].add_run(str(val))
            _set_run_font(run, size_pt=11, bold=False)

    if not relatives:
        doc.add_paragraph()
        _add_paragraph(doc, "Yaqin qarindoshlar kiritilmadi.", size=11)

    # ---------- TELEFON ----------
    phone_p = doc.add_paragraph()
    phone_p.paragraph_format.space_before = Pt(10)
    phone_run = phone_p.add_run(f"Tel: {data.get('phone', '')}")
    _set_run_font(phone_run, size_pt=11, bold=False)
    phone_run.underline = True

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out.read()


DEFAULT_COMMANDS = [
    BotCommand(command="start", description="Botni ishga tushirish"),
    BotCommand(command="menu", description="☰ Bosh menyu"),
    BotCommand(command="malumotnoma", description="Ma'lumotnoma (obyektivka) tuzish"),
    BotCommand(command="namuna", description="Ma'lumotnoma namunasini ko'rish"),
    BotCommand(command="holatim", description="📊 So'rovim holati"),
    BotCommand(command="yordam", description="❓ Yordam"),
    BotCommand(command="hemis", description="🎓 HEMIS hisobini ulash"),
    BotCommand(command="baholarim", description="📊 HEMIS: baho va GPA"),
    BotCommand(command="hemis_uzish", description="🔌 HEMIS hisobini uzish"),
    BotCommand(command="bekor", description="Joriy jarayonni bekor qilish"),
]

ADMIN_ONLY_COMMANDS = [
    BotCommand(command="statistika", description="📊 Bot statistikasi (admin)"),
    BotCommand(command="admin", description="⚙️ Admin panel"),
    BotCommand(command="xabar", description="📢 Hammaga xabar yuborish (admin)"),
    BotCommand(command="adminlar", description="👥 Adminlar ro'yxati (admin)"),
    BotCommand(command="admin_qoshish", description="➕ Admin qo'shish (admin)"),
    BotCommand(command="admin_ochirish", description="➖ Adminni olib tashlash (admin)"),
]


async def _set_admin_commands_for(admin_id: int) -> None:
    """Bitta foydalanuvchi uchun admin buyruqlar menyusini o'rnatadi
    (yangi admin qo'shilganda darhol chaqiriladi)."""
    try:
        await bot.set_my_commands(
            DEFAULT_COMMANDS + ADMIN_ONLY_COMMANDS,
            scope=BotCommandScopeChat(chat_id=admin_id),
        )
    except Exception:
        logging.exception("Admin (%s) uchun buyruqlarni o'rnatishda xatolik", admin_id)


async def _set_commands():
    await bot.set_my_commands(DEFAULT_COMMANDS, scope=BotCommandScopeDefault())
    for admin_id in ADMIN_IDS | db_get_dynamic_admin_ids():
        await _set_admin_commands_for(admin_id)


async def main():
    db_init()
    await _set_commands()
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
